# schedule_web.py

import os
import calendar
import psycopg2
import pandas as pd

from opencc import OpenCC
from dotenv import load_dotenv
from db import get_db, get_conn, db_query
from supabase import create_client
from openpyxl import load_workbook
from psycopg2.extras import RealDictCursor
from sqlalchemy import create_engine, text
from utils import apply_branch_prefix
from schedule.blueprint import schedule_bp
import schedule.routes

from schedule.services.quote_service import get_daily_dharma
from schedule.builders.time_utils import malaysia_today, malaysia_now
from datetime import datetime, timedelta, date, timezone
from schedule.builders.schedule_builder import sync_schedule_after_signup_change
from schedule.services.admin_dashboard_service import load_admin_dashboard_data
from lunar_rules import get_special_day_info, get_next_day_remove_info
from flask import request, session, redirect, url_for, render_template_string, flash
from schedule.builders.schedule_builder import (
    run_schedule_for_date,
    patch_schedule_for_date,
    parse_signup_line_multi,
    normalize_time_text,
    build_buddhist_festival_message,
    get_special_day_info,
    get_next_day_remove_info,
    build_lunar_1_15_message,
    build_normal_message,
    get_duty_targets,
)

from schedule.helpers import (
    get_daily_buddha_quote,
    to_simple,
    normalize_vol_id_for_search,
    time_to_minutes,
    find_volunteer_by_keyword,
    build_monthly_signup_text
)

from schedule.services.settings_service import (
    get_schedule_setting,
    get_schedule_settings,
    save_schedule_setting,
    set_schedule_setting,
)

from schedule.services.publish_service import (
    clear_schedule_need_republish,
    clear_schedule_need_republish,
    is_schedule_published,
    mark_schedule_need_republish,
    publish_schedule_for_date,
    unpublish_schedule_for_date,
    get_schedule_republish_info,
)

from schedule.services.assignment_service import (
    load_assigned_places_for_date,
    load_display_records,
    load_schedule_admin_dashboard_data,
)

from schedule.services.shortage_service import (
    build_shortage_notice_from_assignments,
    build_shortage_summary_for_admin,
    build_signup_shortage_notice
)

from schedule.services.supply_service import (
    load_supply_signups_for_date,
    load_upcoming_supply_signup_alerts,
    load_day_flags,
)

from schedule.constants import (
    PREBOOK_FILE,
    TIME_OPTIONS,
    ROLE_OPTIONS,
)

from schedule.services.whatsapp_service import (
    build_whatsapp_from_assigned
)


load_dotenv()

cc = OpenCC('t2s')  # 繁 → 简

DATABASE_URL = os.environ.get("DATABASE_URL")
engine = create_engine(DATABASE_URL, pool_pre_ping=True) if DATABASE_URL else None

DATABASE_URL = os.environ.get("DATABASE_URL")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATA_DIR = os.path.join(BASE_DIR, "data")

PREBOOK_FILE = os.path.join(
    DATA_DIR,
    "prebook_schedule.xlsx"
)

SCHEDULE_PIN = "1234"

schedule_records = []

def get_supabase_client():
    return create_client(
        os.getenv("SUPABASE_URL"),
        os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    )


def display_report_name(filename):
    name = filename.replace(".xlsx", "")

    if name == "2026_year_report":
        return "2026 年报"

    if "_month_report" in name:
        parts = name.split("_")
        year = parts[0]
        month = parts[1]
        return f"{year}-{month} 月报"

    return filename


def get_pending_signup_counts():
    
    today = malaysia_today()

    tomorrow = today + timedelta(days=1)
    month_start = today.replace(day=1)

    row = db_query("""
        select
            count(*) filter (
                where signup_date = %s
            ) as today_count,

            count(*) filter (
                where signup_date = %s
            ) as tomorrow_count,

            count(*) filter (
                where signup_date >= %s
            ) as month_count

        from volunteer_schedule_signups
        where coalesce(status, 'pending') = 'pending'
        and signup_date >= %s
    """, (
        today,
        tomorrow,
        month_start,
        month_start
    ), fetchone=True)

    return {
        "today": row["today_count"] or 0,
        "tomorrow": row["tomorrow_count"] or 0,
        "month": row["month_count"] or 0,
    }


def search_volunteer_names(keyword, volunteers):
    keyword_simple = to_simple(keyword)

    if not keyword_simple:
        return []

    results = []

    for v in volunteers:
        name = str(v.get("name") or v.get("姓名") or "").strip()
        phone = str(v.get("phone") or v.get("电话") or v.get("电话号码") or "").strip()

        name_simple = to_simple(name)

        if keyword_simple in name_simple or keyword_simple in phone:
            results.append(v)

    return results

def find_name_by_id(vol_id):
    raw_id = str(vol_id or "").strip()
    if not raw_id:
        return raw_id

    ids = [raw_id]

    if "-" in raw_id:
        branch, num = raw_id.split("-", 1)
        branch = branch.strip().upper()
        num = num.strip()
        ids.append(num)

        if branch == "STW" and num.isdigit():
            ids.append("0" + num)
    else:
        if raw_id.startswith("0") and raw_id[1:].isdigit():
            ids.append(f"STW-{raw_id[1:]}")
        else:
            ids.append(f"CHE-{raw_id}")

    ids = list(dict.fromkeys(ids))

    if not engine:
        return raw_id

    try:
        placeholders = ", ".join([f":id{i}" for i in range(len(ids))])
        params = {f"id{i}": v for i, v in enumerate(ids)}

        sql = text(f"""
            SELECT id, name, status, branch
            FROM volunteers
            WHERE id IN ({placeholders})
            LIMIT 1
        """)

        with engine.connect() as conn:
            row = conn.execute(sql, params).mappings().first()

        if not row:
            return raw_id

        return str(row["name"]).strip()

    except Exception as e:
        print("find_name_by_id error:", e)
        return raw_id
    
from opencc import OpenCC
cc = OpenCC("t2s")


@schedule_bp.route("/schedule/logout")
def schedule_logout():

    session.pop("schedule_login", None)

    return redirect("/volunteer")


@schedule_bp.route(
    "/schedule/signup_cancel/<int:signup_id>",
    methods=["POST"]
)
def schedule_signup_cancel(signup_id):

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    with get_conn() as conn:
        with conn.cursor() as cur:

            cur.execute("""
                update volunteer_schedule_signups
                set
                    status = 'cancelled',
                    assigned_place = null,
                    remarks = '负责人取消报名'
                where id = %s
            """, (signup_id,))

            conn.commit()

    sync_schedule_after_signup_change(
        signup_id,
        action="cancel",
        changed_by="admin"
    )

    return redirect("/schedule/signups")

    
def load_buddha_name_options():
    file = os.path.join(BASE_DIR, "data", "fixed_schedule.xlsx")

    if not os.path.exists(file):
        return []

    try:
        df = pd.read_excel(file, sheet_name="佛台固定")
        df.columns = df.columns.astype(str).str.strip()

        names = []
        for col in ["姓名1", "姓名2", "姓名3"]:
            if col in df.columns:
                for v in df[col].dropna():
                    name = str(v).strip()
                    if name and name != "nan" and name not in names:
                        names.append(name)

        return names
    except Exception as e:
        print("load_buddha_name_options error:", e)
        return []

def get_fixed_buddha_for_date(date_str):
    file = os.path.join(BASE_DIR, "data", "fixed_schedule.xlsx")
    print("fixed_schedule file =", file, os.path.exists(file))

    if not os.path.exists(file):
        return []

    try:
        date_obj = pd.to_datetime(date_str)
        weekday_map = {
            0: "星期一",
            1: "星期二",
            2: "星期三",
            3: "星期四",
            4: "星期五",
            5: "星期六",
            6: "星期日",
        }
        weekday = weekday_map[date_obj.weekday()]

        df = pd.read_excel(file, sheet_name="佛台固定")
        df.columns = df.columns.astype(str).str.strip()

        row = df[df["星期"].astype(str).str.strip() == weekday]

        if row.empty:
            return []

        names = []
        r = row.iloc[0]

        for col in ["姓名1", "姓名2", "姓名3"]:
            if col in df.columns:
                name = str(r.get(col, "")).strip()
                if name and name != "nan":
                    names.append(name)

        return names

    except Exception as e:
        print("get_fixed_buddha_for_date error:", e)
        return []

def get_default_time_by_role(role, start_time, end_time):

    role = str(role).strip()

    if role == "值班":
        return start_time, end_time

    if role in ["卫生", "供台", "整理佛台"]:
        return None, None

    return start_time, end_time


def save_prebook_record(record):
    new_df = pd.DataFrame([record])

    if os.path.exists(PREBOOK_FILE):
        try:
            old_df = pd.read_excel(PREBOOK_FILE, sheet_name="预报名")
            old_df.columns = old_df.columns.astype(str).str.strip()
            df = pd.concat([old_df, new_df], ignore_index=True)
        except Exception:
            df = new_df
    else:
        df = new_df

    df = df.drop_duplicates(
        subset=["日期", "姓名", "岗位", "开始时间", "结束时间"],
        keep="first"
    )

    with pd.ExcelWriter(PREBOOK_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="预报名", index=False)

    
def get_signup_summary_for_date(date_str):
    from db import db_query

    row = db_query("""
        select
            count(*) as total,
            count(*) filter (
                where coalesce(status, 'pending') = 'pending'
            ) as pending,
            count(*) filter (
                where status = 'assigned'
            ) as assigned,
            count(*) filter (
                where status = 'cancelled'
            ) as cancelled
        from volunteer_schedule_signups
        where signup_date = %s
    """, (date_str,), fetchone=True)

    return {
        "total": row["total"] or 0,
        "pending": row["pending"] or 0,
        "assigned": row["assigned"] or 0,
        "cancelled": row["cancelled"] or 0,
    }
    

def save_buddha_override(date_str, final_names):
    file = os.path.join(BASE_DIR, "buddha_override.xlsx")

    names = [str(n).strip() for n in final_names if str(n).strip()]

    new_row = {
        "日期": date_str,
        "姓名1": names[0] if len(names) > 0 else "",
        "姓名2": names[1] if len(names) > 1 else "",
        "姓名3": names[2] if len(names) > 2 else "",
    }

    if os.path.exists(file):
        df = pd.read_excel(file)
        df.columns = df.columns.astype(str).str.strip()
        df = df[df["日期"].astype(str) != date_str]
        df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
    else:
        df = pd.DataFrame([new_row])

    with pd.ExcelWriter(file, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)

def delete_prebook_record(record):
    if not os.path.exists(PREBOOK_FILE):
        return

    try:
        df = pd.read_excel(PREBOOK_FILE, sheet_name="预报名")
        df.columns = df.columns.astype(str).str.strip()

        cond = (
            (df["日期"].astype(str) == str(record["日期"])) &
            (df["姓名"].astype(str) == str(record["姓名"])) &
            (df["岗位"].astype(str) == str(record["岗位"])) &
            (df["开始时间"].astype(str) == str(record["开始时间"])) &
            (df["结束时间"].astype(str) == str(record["结束时间"]))
        )

        df = df[~cond]

        with pd.ExcelWriter(PREBOOK_FILE, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="预报名", index=False)

    except Exception as e:
        print("delete_prebook_record error:", e)


@schedule_bp.route("/schedule/copy_whatsapp", methods=["POST"])
def schedule_copy_whatsapp():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    date_str = request.form.get("date", "").strip()

    if not date_str:
        return "❌ 没有日期<br><a href='/schedule?mode=day'>返回</a>"

    output = build_whatsapp_from_assigned(date_str)

    clear_schedule_need_republish(date_str)

    return render_template_string(DAY_OUTPUT_HTML, output=output)


def is_big_day(date_str):
    from datetime import datetime

    date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
    info = get_special_day_info(date_obj)

    return info.get("template_type") in [
        "lunar_1_15",
        "buddhist_festival"
    ]


@schedule_bp.route(
    "/schedule/signup_restore/<int:signup_id>",
    methods=["POST"]
)
def schedule_signup_restore(signup_id):

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    with get_conn() as conn:
        with conn.cursor() as cur:

            cur.execute("""
                update volunteer_schedule_signups
                set
                    status = 'pending',
                    assigned_place = null,
                    remarks = '负责人恢复报名'
                where id = %s
            """, (signup_id,))

            conn.commit()

    sync_schedule_after_signup_change(
        signup_id,
        action="upsert",
        changed_by="admin"
    )

    return redirect("/schedule/signups")


def build_month_prebook_message(year, month):
    import calendar
    from datetime import date

    first_day = date(year, month, 1)
    last_day = date(year, month, calendar.monthrange(year, month)[1])

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                select
                    signup_date,
                    name,
                    role,
                    start_time,
                    end_time
                from volunteer_schedule_signups
                where signup_date between %s and %s
                and coalesce(status, 'pending') <> 'cancelled'
                order by signup_date, role, start_time, name
            """, (first_day, last_day))
            rows = cur.fetchall()

    by_day = {}
    for r in rows:
        d = r["signup_date"]
        by_day.setdefault(d, []).append(r)

    weekday_map = {
        0: "星期一",
        1: "星期二",
        2: "星期三",
        3: "星期四",
        4: "星期五",
        5: "星期六",
        6: "星期日",
    }

    lines = []
    lines.append(f"{year}年{month}月份预报名名单")
    lines.append("")
    lines.append("以下为目前已收到的预报名资料，最终岗位安排以负责人公布值班表为准。")
    lines.append("")

    days_in_month = calendar.monthrange(year, month)[1]

    for day in range(1, days_in_month + 1):
        d = date(year, month, day)
        day_rows = by_day.get(d, [])

        if not day_rows:
            continue

        cleaning = []
        supply = []
        duty = []
        others = []

        for r in day_rows:
            name = str(r["name"]).strip()
            role = str(r["role"]).strip()
            start = r.get("start_time") or ""
            end = r.get("end_time") or ""

            if role == "卫生":
                cleaning.append(name)

            elif role == "供台":
                supply.append(name)

            elif role == "值班":
                if start and end:
                    duty.append(f"{name} {start}~{end}")
                else:
                    duty.append(name)

            else:
                others.append(f"{name}（{role}）")

        lines.append(f"{day}/{month}/{year} {weekday_map[d.weekday()]}")

        if cleaning:
            lines.append("卫生：" + "、".join(cleaning))

        if supply:
            lines.append("供台：" + "、".join(supply))

        if duty:
            lines.append("值班：")
            lines.extend(duty)

        if others:
            lines.append("其他：")
            lines.extend(others)

        lines.append("")

    return "\n".join(lines)


@schedule_bp.route("/schedule/edit_assigned/<int:assignment_id>", methods=["GET", "POST"])
def edit_assigned_place(assignment_id):

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    shift_options = ["绿班", "橙班", "黄班"]
    place_options = [
        "观音堂",
        "活动中心",
        "佛堂卫生",
        "二楼卫生",
        "楼梯卫生",
        "设师父供台",
        "整理佛台",
    ]

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:

            if request.method == "POST":
                shift_label = request.form.get("shift_label", "").strip()
                assigned_place = request.form.get("assigned_place", "").strip()
                start_time = request.form.get("start_time", "").strip()
                end_time = request.form.get("end_time", "").strip()

                cur.execute("""
                    select assignment_date
                    from volunteer_schedule_assignments
                    where id = %s
                """, (assignment_id,))

                assignment_info = cur.fetchone()

                if not assignment_info:
                    return "❌ 找不到排班记录<br><a href='/schedule/signups'>返回</a>"

                assignment_date = assignment_info["assignment_date"]

                cur.execute("""
                    update volunteer_schedule_assignments
                    set
                        shift_label = %s,
                        assigned_place = %s,
                        start_time = %s,
                        end_time = %s,
                        remarks = '负责人调整排班',
                        assignment_source = 'admin',
                        locked_by_admin = true
                    where id = %s
                """, (
                    shift_label or None,
                    assigned_place,
                    start_time or None,
                    end_time or None,
                    assignment_id
                ))

                conn.commit()

                mark_schedule_need_republish(assignment_date)

                return redirect("/schedule/signups")

            cur.execute("""
                select *
                from volunteer_schedule_assignments
                where id = %s
            """, (assignment_id,))

            row = cur.fetchone()

    if not row:
        return "❌ 找不到排班记录<br><a href='/schedule/signups'>返回</a>"

    return render_template_string("""
<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>调整排班</title>
<style>
body { font-family:"Microsoft YaHei", Arial; background:#f5f5f5; padding:20px; }
.box { background:white; max-width:700px; margin:auto; padding:25px; border-radius:15px; }
input, select, button { font-size:24px; padding:10px; margin:8px 0; width:100%; box-sizing:border-box; }
p { font-size:22px; }
</style>
</head>
<body>
<div class="box">

<h1>✏️ 调整排班</h1>

<p><b>{{ row.name }}</b>（{{ row.volunteer_id or "" }}）</p>
<p>日期：{{ row.assignment_date }}</p>
<p>岗位：{{ row.role }}</p>

{% if row.locked_by_admin %}
    <div style="background:#fff3cd;border:2px solid #f0c36d;border-radius:12px;padding:14px;font-size:22px;margin:15px 0;">
        🔒 已由负责人锁定<br>
        <small>系统不会自动覆盖这笔安排</small>
    </div>
{% else %}
    <div style="background:#e8f5e9;border:2px solid #81c784;border-radius:12px;padding:14px;font-size:22px;margin:15px 0;">
        🤖 系统自动安排<br>
        <small>保存调整后将变成负责人锁定</small>
    </div>
{% endif %}

<form method="post">

    <label>班别</label>
    <select name="shift_label">
        <option value="">无</option>
        {% for s in shift_options %}
        <option value="{{ s }}" {% if row.shift_label == s %}selected{% endif %}>
            {{ s }}
        </option>
        {% endfor %}
    </select>

    <label>地点</label>
    <select name="assigned_place" required>
        {% for p in place_options %}
        <option value="{{ p }}" {% if row.assigned_place == p %}selected{% endif %}>
            {{ p }}
        </option>
        {% endfor %}
    </select>

    <label>开始时间</label>
    <input name="start_time" value="{{ row.start_time or '' }}" placeholder="例如 11:00am">

    <label>结束时间</label>
    <input name="end_time" value="{{ row.end_time or '' }}" placeholder="例如 2:00pm">

    <button type="submit">💾 保存调整</button>

</form>

<br>
<a href="/schedule/signups">⬅ 返回报名管理</a>

</div>
</body>
</html>
""",
        row=row,
        shift_options=shift_options,
        place_options=place_options
    )


def auto_patch_new_signup_to_assignment(signup_id):
    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:

            cur.execute("""
                select *
                from volunteer_schedule_signups
                where id = %s
                and coalesce(status, 'pending') <> 'cancelled'
            """, (signup_id,))
            s = cur.fetchone()

            if not s:
                return

            cur.execute("""
                select id
                from volunteer_schedule_assignments
                where signup_id = %s
                and coalesce(status, 'assigned') <> 'cancelled'
                limit 1
            """, (signup_id,))
            exists = cur.fetchone()

            if exists:
                return

            role = s["role"]
            signup_date = s["signup_date"]
            name = s["name"]
            volunteer_id = s["volunteer_id"]
            start_time = s["start_time"]
            end_time = s["end_time"]

            if role == "供台":
                return

            if role == "卫生":
                assigned_place = "佛堂卫生"
                shift_label = "卫生"

            else:
                s_min = time_to_minutes(start_time)
                e_min = time_to_minutes(end_time)

                if s_min is None or e_min is None:
                    return

                if s_min < 14 * 60:
                    shift_label = "橙"
                else:
                    shift_label = "黄"

                cur.execute("""
                    select assigned_place, count(*) as cnt
                    from volunteer_schedule_assignments
                    where assignment_date = %s
                    and shift_label = %s
                    and assigned_place in ('观音堂', '活动中心')
                    and coalesce(status, 'assigned') <> 'cancelled'
                    group by assigned_place
                """, (signup_date, shift_label))

                count_rows = cur.fetchall()
                counts = {"观音堂": 0, "活动中心": 0}

                for r in count_rows:
                    counts[r["assigned_place"]] = r["cnt"]

                if counts["观音堂"] <= counts["活动中心"]:
                    assigned_place = "观音堂"
                else:
                    assigned_place = "活动中心"

            cur.execute("""
                insert into volunteer_schedule_assignments
                (signup_id, volunteer_id, name, assignment_date, role,
                 shift_label, assigned_place, start_time, end_time, status, remarks)
                values (%s, %s, %s, %s, %s,
                        %s, %s, %s, %s, 'assigned', '报名后自动补排')
            """, (
                s["id"],
                volunteer_id,
                name,
                signup_date,
                role,
                shift_label,
                assigned_place,
                start_time,
                end_time
            ))

            cur.execute("""
                update volunteer_schedule_signups
                set status = 'assigned'
                where id = %s
            """, (s["id"],))

            conn.commit()


def build_shortage_notice(date_str):

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:

            cur.execute("""
                select role
                from volunteer_schedule_signups
                where signup_date = %s
                and coalesce(status,'pending') <> 'cancelled'
            """, (date_str,))

            rows = cur.fetchall()

    cleaning_count = 0
    duty_count = 0

    for r in rows:

        role = str(r.get("role") or "").strip()

        if role == "卫生":
            cleaning_count += 1

        elif role == "值班":
            duty_count += 1

    notices = []

    if cleaning_count < 2:
        notices.append(
            f"卫生可以多加 {2-cleaning_count} 位义工。"
        )

    if duty_count < 6:
        notices.append(
            "全日值班可以多加 2~4 位义工。"
        )

    if not notices:
        return f"""
师兄们，大家好！

{date_str} 的报名人数暂时充足。

感恩大家发心护持观音堂。
🙏🙏🙏
"""

    body = "\n".join(notices)

    return f"""
师兄们，大家好！

{body}

感恩大家🙏🙏🙏
"""

def build_signup_notice_whatsapp(date_str):

    date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()

    weekday_map = {
        0: "星期一",
        1: "星期二",
        2: "星期三",
        3: "星期四",
        4: "星期五",
        5: "星期六",
        6: "星期日",
    }

    date_text = f"{date_obj.day}/{date_obj.month}/{date_obj.year} ({weekday_map[date_obj.weekday()]})"

    fixed_buddha = get_fixed_buddha_for_date(date_str)

    groups = {
        "佛堂卫生": [],
        "二楼卫生": [],
        "楼梯卫生": [],
        "橙观音堂": [],
        "橙活动中心": [],
        "黄观音堂": [],
        "黄活动中心": [],
    }

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                select *
                from volunteer_schedule_assignments
                where assignment_date = %s
                and coalesce(status, 'assigned') <> 'cancelled'
                order by start_time, name
            """, (date_str,))
            rows = cur.fetchall()

    for r in rows:
        place = r.get("assigned_place")
        shift = r.get("shift_label")
        key = None

        if place in ["佛堂卫生", "二楼卫生", "楼梯卫生"]:
            key = place
        elif shift and place:
            key = f"{shift}{place}"

        if key in groups:
            groups[key].append(r)

    def format_people(rows):
        if not rows:
            return ""

        lines = []
        for r in rows:
            name = r.get("name") or ""
            start = r.get("start_time") or ""
            end = r.get("end_time") or ""

            lines.append(str(name))

            if start and end:
                lines.append(f"{start}~{end}")

        return "\n".join(lines)

    buddha_text = "　".join(fixed_buddha) if fixed_buddha else ""

    text = f"""师兄们大家好！

{date_text}

十点正请安

8:00am~10:00am  或      
8:00am~完成佛台工作 
整理佛台: 
{buddha_text}


8:00am~10:00am 或 
8:00am~完成卫生工作 
佛堂卫生: 
{format_people(groups["佛堂卫生"])}
二楼卫生: 
{format_people(groups["二楼卫生"])}
楼梯卫生: 
{format_people(groups["楼梯卫生"])}
每日卫生义工请注意：清理完卫生之后，请把卫生用具包括吸尘机放回原位（活动中心store里面的小房间）



10:00am~2:00pm 
🟠 观音堂: 
{format_people(groups["橙观音堂"])}
🟠 活动中心: 
{format_people(groups["橙活动中心"])}

2:00pm~6:00pm 
🟡 观音堂: 
{format_people(groups["黄观音堂"])}

🟡 活动中心: 
{format_people(groups["黄活动中心"])}


观音堂早晚香 由值班义工带领上香。

观音堂续香
佛友可以续香（黑色无烟香），但是要跟着请安词，必须燃烧完了一支香才续香。如有多位佛友要续香，请先让给先到达观音堂的佛友，轮流续香。

值班义工请注意
1）下雨天记得关上烧送小房子的窗口。
2）在离开观音堂之前，请确保把所有的窗口关上。
3）观音堂第一架的冷气坚决不能调。第二架和第三架可以轮流。

另外，请大家多留意义工群信息，以便大家能够团结一致的护持好观音堂。佛子齐心，普度众生。

义工报名请点击以下链接：
https://gyt-checkin.onrender.com/volunteer

非常感恩大家！
大家功德无量！
🙏🙏🙏
"""

    return text


@schedule_bp.route("/schedule", methods=["GET", "POST"])
@schedule_bp.route("/schedule/admin", methods=["GET", "POST"])
def schedule_admin():
    import time

    if not session.get("schedule_login"):
        if request.method == "POST":
            pin = request.form.get("pin", "").strip()

            if pin == SCHEDULE_PIN:
                session.permanent = True
                session["schedule_login"] = True
                return redirect(url_for("schedule.schedule_admin"))

            return "❌ PIN 错误<br><a href='/schedule'>返回</a>"

        return render_template_string(LOGIN_HTML)

    t0 = time.time()

    now = malaysia_now()

    switch_time = get_schedule_setting(
        "default_day_switch_time",
        "18:00"
    )

    try:
        switch_hour, switch_minute = map(int, switch_time.split(":"))
    except:
        switch_hour, switch_minute = 18, 0

    change_time = now.replace(
        hour=switch_hour,
        minute=switch_minute,
        second=0,
        microsecond=0
    )

    if now >= change_time:
        default_schedule_date = (now + timedelta(days=1)).strftime("%Y-%m-%d")
    else:
        default_schedule_date = now.strftime("%Y-%m-%d")

    mode = request.args.get("mode", "")

    override_date = request.args.get("override_date") or default_schedule_date
    t1 = time.time()
    dashboard = load_admin_dashboard_data(
        mode=mode,
        override_date=override_date,
    )
    print("dashboard:", round(time.time() - t1, 2))
    print("NOW =", now)
    print("SWITCH =", switch_time)
    print("DEFAULT =", default_schedule_date)
    print("OVERRIDE =", override_date)

    t2 = time.time()
    buddha_names = load_buddha_name_options()
    print("buddha_names:", round(time.time() - t2, 2))

    t3 = time.time()
    fixed_buddha_today = get_fixed_buddha_for_date(override_date)
    print("fixed_buddha_today:", round(time.time() - t3, 2))

    print("schedule_admin total:", round(time.time() - t0, 2))

    return render_template_string(
        SCHEDULE_HTML,
        mode=mode,
        times=TIME_OPTIONS,
        roles=ROLE_OPTIONS,
        tomorrow=default_schedule_date,
        buddha_names=buddha_names,
        override_date=override_date,
        fixed_buddha_today=fixed_buddha_today,
        multi_day_signup_open=dashboard["multi_day_signup_open"],
        meal_signup_open=dashboard["meal_signup_open"],
        records=dashboard["records"],
        pending_counts=dashboard["pending_counts"],
        day_summary=dashboard["day_summary"],
        is_today_or_past=dashboard["is_today_or_past"],
        day_flags=dashboard["day_flags"],
        special_day_info=dashboard["special_day_info"],
        remove_info=dashboard["remove_info"],
        shortage_summary=dashboard["shortage_summary"],
        supply_signups=dashboard["supply_signups"],
        supply_alerts=dashboard["supply_alerts"],
        is_published=dashboard["is_published"],
        republish_info=dashboard["republish_info"],
        default_year=dashboard["default_year"],
        default_month=dashboard["default_month"],
    )


# ============================================================
# ⚠️ LEGACY：旧版负责人录入入口
#
# 这支 route 属于旧流程：
# - 使用 schedule_records
# - 使用 prebook_schedule.xlsx
# - 使用 save_prebook_record()
# - 使用旧版 monthly_prebook_message
#
# 新流程已经改为：
# signups → assignments → publish → patch
#
# 以后不要再把新功能写进这里。
# 如需新增功能，请另开新的 route / service。
# ============================================================
@schedule_bp.route("/schedule/add", methods=["POST"])
def schedule_add():
    print("⚠️ LEGACY route used: /schedule/add")
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    mode = request.form.get("mode", "").strip()

    vol_keyword = request.form.get("keyword", "").strip()

    branch = request.form.get("branch", "CHE").strip().upper()
    vol_keyword = apply_branch_prefix(vol_keyword, branch)
    roles = request.form.getlist("roles")
    start_time = request.form.get("start_time", "").strip()
    end_time = request.form.get("end_time", "").strip()

    action = request.form.get("action", "add")
    single_date = request.form.get("single_date", "").strip()
    year = request.form.get("year", "").strip()
    month = request.form.get("month", "").strip()
    days = request.form.getlist("days")

    if mode == "prebook" and action == "generate_monthly":
        return """
        <h1>⚠️ 旧版月预报名功能已停用</h1>
        <p>请使用新的「多日报名」和「查看预报名名单」。</p>
        <a href="/schedule?mode=prebook">返回</a>
        """
        year = request.form.get("year", "").strip()
        month = request.form.get("month", "").strip()

        try:
            output = generate_monthly_prebook_message(int(year), int(month))
        except Exception as e:
            output = f"❌ 生成失败：{e}"

        return render_template_string(
            MONTHLY_PREBOOK_HTML,
            output=output
        )

    if mode == "day" and action == "parse_raw":
        if not single_date:
            single_date = request.form.get("date", "").strip()

        if not single_date:
            return "❌ 请先选择日期<br><a href='/schedule?mode=day'>返回</a>"

        raw_signup = request.form.get("raw_signup", "").strip()

        if not raw_signup:
            return "❌ 请先贴上 WhatsApp 报名内容<br><a href='/schedule?mode=day'>返回</a>"

        lines = raw_signup.splitlines()

        for line in lines:
            line = normalize_time_text(line.strip())
            if not line:
                continue

            parsed_list = parse_signup_line_multi(line)

            for p in parsed_list:
                name = str(p.get("姓名", "")).strip()
                role = str(p.get("岗位", "值班")).strip()

                job_start = p.get("开始时间") or "10:00am"
                job_end = p.get("结束时间") or "2:00pm"

                record = {
                    "日期": single_date,
                    "编号": "",
                    "姓名": name,
                    "岗位": role,
                    "开始时间": job_start,
                    "结束时间": job_end,
                    "备注": "WhatsApp解析",
                }

                schedule_records.append(record)
                save_prebook_record(record)

        return redirect(url_for("schedule.schedule_admin", mode="day"))

    if not vol_keyword:
        return "❌ 请填写义工编号<br><a href='/schedule/admin'>返回</a>"

    if not roles:
        return "❌ 请至少选择一个岗位<br><a href='/schedule/admin'>返回</a>"

    matches = find_volunteer_by_keyword(vol_keyword)

    if not matches:
        return "❌ 找不到义工<br><a href='/schedule'>返回</a>"

    if len(matches) > 1:
        return render_template_string("""
        <h3>找到多个义工，请选择：</h3>

        {% for v in matches %}
            <form method="post" action="/schedule/prebook_add">
                <input type="hidden" name="vol_id" value="{{ v.id }}">
                <input type="hidden" name="mode" value="{{ mode }}">
                <input type="hidden" name="action" value="{{ action }}">

                <input type="hidden" name="year" value="{{ year }}">
                <input type="hidden" name="month" value="{{ month }}">

                {% for d in days %}
                    <input type="hidden" name="days" value="{{ d }}">
                {% endfor %}

                {% for r in roles %}
                    <input type="hidden" name="roles" value="{{ r }}">
                {% endfor %}

                <input type="hidden" name="start_time" value="{{ start_time }}">
                <input type="hidden" name="end_time" value="{{ end_time }}">
                <input type="hidden" name="single_date" value="{{ single_date }}">

                <button type="submit">{{ v.name }} ({{ v.id }})</button>
            </form>
        {% endfor %}
        """,
        matches=matches,
        mode=mode,
        action=action,
        roles=roles,
        start_time=start_time,
        end_time=end_time,
        single_date=single_date,
        year=year,
        month=month,
        days=days
        )

    vol = matches[0]
    vol_id = str(vol["id"])
    name = str(vol["name"])

    if mode == "day":
        if not single_date:
            single_date = request.form.get("date", "").strip()

        if not single_date:
            return "❌ 请先选择日期<br><a href='/schedule?mode=day'>返回</a>"

        date_list = [single_date]

    elif mode == "prebook":
        if not year or not month:
            return "❌ 请选择年份和月份<br><a href='/schedule?mode=prebook'>返回</a>"

        if action == "generate_monthly":
            return """
            <h1>⚠️ 旧版月预报名功能已停用</h1>
            <p>请使用新的「多日报名」和「查看预报名名单」。</p>
            <a href="/schedule?mode=prebook">返回</a>
            """
            try:
                output = generate_monthly_prebook_message(
                    int(year),
                    int(month)
                )
            except Exception as e:
                output = f"❌ 生成失败：{e}"

            return render_template_string(
                MONTHLY_PREBOOK_HTML,
                output=output
            )

        if not days:
            return "❌ 请至少选择一天<br><a href='/schedule?mode=prebook'>返回</a>"

        month_full = f"{int(year)}-{int(month):02d}"
        date_list = [
            f"{month_full}-{int(day):02d}"
            for day in days
        ]

    else:
        return "❌ 模式错误<br><a href='/schedule'>返回</a>"

    for date_text in date_list:
        for role in roles:
            job_start, job_end = get_default_time_by_role(role, start_time, end_time)

            record = {
                "日期": date_text,
                "编号": vol_id,
                "姓名": name,
                "岗位": role,
                "开始时间": job_start,
                "结束时间": job_end,
                "备注": "网页录入",
            }

            schedule_records.append(record)
            save_prebook_record(record)

    return redirect(url_for("schedule.schedule_admin", mode=mode))


@schedule_bp.route("/schedule/toggle_setting", methods=["POST"])
def toggle_schedule_setting():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    key = request.form.get("key", "").strip()
    force = request.form.get("force", "").strip()
    return_to = request.form.get("return_to") or url_for("schedule.schedule_admin")

    allowed_keys = {
        "multi_day_signup_open",
        "meal_signup_open",
    }

    if key not in allowed_keys:
        return "❌ 不允许的设置"

    if force in ("true", "false"):
        set_schedule_setting(key, force, updated_by="admin")
    else:
        current = get_schedule_setting(key, "false")
        new_value = "false" if current == "true" else "true"
        set_schedule_setting(key, new_value, updated_by="admin")

    return redirect(return_to)


@schedule_bp.route("/schedule/prebook_add", methods=["POST"])
def schedule_prebook_add():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    keyword = request.form.get("vol_id", "").strip()
    single_date = request.form.get("single_date", "").strip()

    roles = request.form.getlist("roles")

    start_time = request.form.get("start_time", "").strip()
    end_time = request.form.get("end_time", "").strip()

    if not keyword:
        return "❌ 请输入义工编号 / 姓名<br><a href='/schedule?mode=day'>返回</a>"

    if not single_date:
        return "❌ 缺少报名日期<br><a href='/schedule?mode=day'>返回</a>"

    if not roles:
        return "❌ 请选择岗位<br><a href='/schedule?mode=day'>返回</a>"

    try:
        signup_date = datetime.strptime(single_date, "%Y-%m-%d").date()
    except ValueError:
        return "❌ 日期格式错误<br><a href='/schedule?mode=day'>返回</a>"

    matches = find_volunteer_by_keyword(keyword)

    if not matches:
        return "❌ 找不到义工<br><a href='/schedule?mode=day'>返回</a>"

    if len(matches) > 1:
        return "❌ 找到多个同名义工，请用义工编号查询<br><a href='/schedule?mode=day'>返回</a>"

    vol = matches[0]
    volunteer_id = str(vol["id"])
    name = str(vol["name"])

    inserted = 0
    skipped = 0

    with get_conn() as conn:
        with conn.cursor() as cur:

            for role in roles:

                role = str(role).strip()

                if role == "值班":
                    role_start = start_time
                    role_end = end_time

                    if not role_start or not role_end:
                        skipped += 1
                        continue

                    s_min = time_to_minutes(role_start)
                    e_min = time_to_minutes(role_end)

                    if s_min is None or e_min is None or e_min <= s_min:
                        skipped += 1
                        continue

                else:
                    role_start = None
                    role_end = None

                cur.execute("""
                    select id
                    from volunteer_schedule_signups
                    where volunteer_id = %s
                    and signup_date = %s
                    and role = %s
                    and coalesce(status, 'pending') <> 'cancelled'
                    limit 1
                """, (
                    volunteer_id,
                    signup_date,
                    role
                ))

                exists = cur.fetchone()

                if exists:
                    skipped += 1
                    continue

                cur.execute("""
                    insert into volunteer_schedule_signups (
                        volunteer_id,
                        name,
                        signup_date,
                        role,
                        start_time,
                        end_time,
                        status,
                        remarks,
                        created_at
                    )
                    values (
                        %s, %s, %s, %s, %s, %s,
                        'pending',
                        '负责人代报名',
                        now()
                    )
                """, (
                    volunteer_id,
                    name,
                    signup_date,
                    role,
                    role_start,
                    role_end
                ))

                inserted += 1

        conn.commit()

    return f"""
    <h1>✅ 负责人代报名已加入</h1>
    <p>日期：{signup_date}</p>
    <p>义工：{name}</p>
    <p>成功加入：{inserted} 笔</p>
    <p>跳过重复/无效：{skipped} 笔</p>
    <a href="/schedule?mode=day&override_date={signup_date}">返回当天安排</a><br>
    <a href="/schedule/admin">返回后台</a>
    """
    
@schedule_bp.route("/schedule/generate_day", methods=["POST"])
def schedule_generate_day():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    date = (
        request.form.get("date")
        or request.form.get("single_date")
        or request.args.get("date")
        or ""
    ).strip()

    if not date:
        return "❌ 没有收到日期，请返回选择日期<br><a href='/schedule?mode=day'>返回</a>"

    if is_schedule_published(date):
        return """
        <h1>🔒 已发布，不能重新自动排班</h1>
        <p>这一天已经发布正式值班表。</p>
        <p>发布后新报名义工会自动补排，不需要重新自动排班。</p>
        <a href="/schedule?mode=day">返回负责人首页</a>
        """

    target_date = datetime.strptime(date, "%Y-%m-%d").date()

    if target_date < malaysia_today():
        return """
        <h1>❌ 当天排班已锁定</h1>
        <p>过去日期不能重新自动排班，只能使用「补入待安排」。</p>
        <a href="/schedule?mode=day">返回</a>
        """

    try:
        output = run_schedule_for_date(date)
    except Exception as e:
        output = f"❌ 生成失败：{e}"

    return render_template_string(DAY_OUTPUT_HTML, output=output)


@schedule_bp.route("/schedule/monthly_prebook")
def schedule_monthly_prebook():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    year = int(request.args.get("year"))
    month = int(request.args.get("month"))

    output = build_month_prebook_message(year, month)

    return render_template_string(
        MONTHLY_PREBOOK_HTML,
        output=output
    )


@schedule_bp.route("/schedule/clear", methods=["POST"])
def schedule_clear():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    schedule_records.clear()
    return redirect(url_for("schedule.schedule_admin"))

@schedule_bp.route("/schedule/delete/<int:index>", methods=["POST"])
def schedule_delete(index):
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    mode = request.form.get("mode", "day")

    if 0 <= index < len(schedule_records):
        record = schedule_records[index]
        delete_prebook_record(record)
        schedule_records.pop(index)

    return redirect(url_for("schedule.schedule_admin", mode=mode))
      

def load_pending_signups_for_date(date_str):
    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                select
                    id,
                    volunteer_id,
                    name,
                    role,
                    start_time,
                    end_time
                from volunteer_schedule_signups
                where signup_date = %s
                and coalesce(status, 'pending') = 'pending'
                order by role, start_time, name
            """, (date_str,))

            return cur.fetchall()
        

def get_recent_cleaning_places(volunteer_id, before_date, limit=5):
    if not volunteer_id:
        return []

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                select assigned_place
                from volunteer_schedule_assignments
                where volunteer_id = %s
                and assignment_date < %s
                and role = '卫生'
                and assigned_place in ('佛堂卫生', '二楼卫生', '楼梯卫生')
                and coalesce(status, 'assigned') <> 'cancelled'
                order by assignment_date desc, id desc
                limit %s
            """, (volunteer_id, before_date, limit))

            rows = cur.fetchall()

    return [r["assigned_place"] for r in rows if r.get("assigned_place")]


def fill_pending_volunteers(date_str):
    assigned_rows = load_assigned_places_for_date(date_str)
    pending_rows = load_pending_signups_for_date(date_str)

    if not pending_rows:
        return "✅ 没有待安排报名"

    place_count = {
        "佛堂卫生": 0,
        "二楼卫生": 0,
        "楼梯卫生": 0,
        "设师父供台": 0,
        "绿观音堂": 0,
        "绿活动中心": 0,
        "橙观音堂": 0,
        "橙活动中心": 0,
        "黄观音堂": 0,
        "黄活动中心": 0,
    }

    for r in assigned_rows:
        shift_label = r.get("shift_label") or ""
        place = r.get("assigned_place") or ""

        key = f"{shift_label}{place}" if shift_label else place

        if key in place_count:
            place_count[key] += 1
        elif place in place_count:
            place_count[place] += 1

    with get_conn() as conn:
        with conn.cursor() as cur:

            for r in pending_rows:
                signup_id = r["id"]
                volunteer_id = r.get("volunteer_id")
                name = r["name"]
                role = r["role"]
                start_time = r.get("start_time")
                end_time = r.get("end_time")

                assigned_place = None
                shift_label = None

                if role == "卫生":
                    cleaning_places = ["佛堂卫生", "二楼卫生", "楼梯卫生"]

                    recent_places = get_recent_cleaning_places(
                        volunteer_id,
                        date_str,
                        limit=5
                    )

                    def cleaning_score(place):
                        score = place_count.get(place, 0) * 10

                        if recent_places:
                            if place == recent_places[0]:
                                score += 100

                            if place in recent_places[:3]:
                                score += 30

                            if place in recent_places:
                                score += 10

                        return score

                    assigned_place = min(
                        cleaning_places,
                        key=cleaning_score
                    )

                elif role == "供台":
                    assigned_place = "设师父供台"

                elif role == "值班":
                    duty_places = [
                        ("绿班", "观音堂", "绿观音堂"),
                        ("绿班", "活动中心", "绿活动中心"),
                        ("橙班", "观音堂", "橙观音堂"),
                        ("橙班", "活动中心", "橙活动中心"),
                        ("黄班", "观音堂", "黄观音堂"),
                        ("黄班", "活动中心", "黄活动中心"),
                    ]

                    shift_label, assigned_place, key = min(
                        duty_places,
                        key=lambda x: place_count[x[2]]
                    )

                else:
                    assigned_place = role

                key = f"{shift_label.replace('班', '')}{assigned_place}" if shift_label else assigned_place
                if key in place_count:
                    place_count[key] += 1

                # 写入真正排班表
                cur.execute("""
                    insert into volunteer_schedule_assignments (
                        signup_id,
                        volunteer_id,
                        name,
                        assignment_date,
                        role,
                        shift_label,
                        assigned_place,
                        start_time,
                        end_time,
                        status,
                        remarks,
                        created_at
                    )
                    values (
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, 'assigned', %s, now()
                    )
                """, (
                    signup_id,
                    volunteer_id,
                    name,
                    date_str,
                    role,
                    shift_label,
                    assigned_place,
                    start_time,
                    end_time,
                    "负责人补入待安排"
                ))

                # 报名表只改状态
                cur.execute("""
                    update volunteer_schedule_signups
                    set
                        status = 'assigned',
                        remarks = '负责人补入待安排'
                    where id = %s
                """, (signup_id,))

            conn.commit()

    return f"✅ 已补入 {len(pending_rows)} 位待安排义工"


@schedule_bp.route("/schedule/reports")
def schedule_reports():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    today = malaysia_today()

    ym = request.args.get("ym", today.strftime("%Y-%m"))

    start_date = f"{ym}-01"

    y, m = map(int, ym.split("-"))
    if m == 12:
        next_month = f"{y+1}-01-01"
    else:
        next_month = f"{y}-{m+1:02d}-01"

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:

            # 本月签到人数
            cur.execute("""
                select count(distinct volunteer_id) as total_people
                from volunteer_attendance_logs
                where check_in_time >= %s
                and check_in_time < %s
            """, (start_date, next_month))
            total_people = cur.fetchone()["total_people"]

            # 本月签到次数
            cur.execute("""
                select count(*) as total_checkins
                from volunteer_attendance_logs
                where check_in_time >= %s
                and check_in_time < %s
            """, (start_date, next_month))
            total_checkins = cur.fetchone()["total_checkins"]

            # 义工排行榜
            cur.execute("""
                select
                    name,
                    count(*) as checkin_count
                from volunteer_attendance_logs
                where check_in_time >= %s
                and check_in_time < %s
                group by name
                order by checkin_count desc, name
                limit 30
            """, (start_date, next_month))
            ranking = cur.fetchall()

            # 岗位统计
            cur.execute("""
                select
                    assigned_place,
                    count(*) as total
                from volunteer_schedule_assignments
                where assignment_date >= %s
                and assignment_date < %s
                and coalesce(status, '') <> 'cancelled'
                group by assigned_place
                order by total desc
            """, (start_date, next_month))
            place_stats = cur.fetchall()

            # 已排班未签到
            cur.execute("""
                select
                    a.assignment_date,
                    a.name,
                    a.role,
                    a.assigned_place
                from volunteer_schedule_assignments a
                left join volunteer_attendance_logs l
                    on l.volunteer_id = a.volunteer_id
                    and date(l.check_in_time) = a.assignment_date
                where a.assignment_date >= %s
                and a.assignment_date < %s
                and coalesce(a.status, '') <> 'cancelled'
                and l.id is null
                order by a.assignment_date, a.name
            """, (start_date, next_month))
            absent_rows = cur.fetchall()

    return render_template_string("""
    <h1>📊 排班报表中心</h1>

    <form method="get">
        <label>月份：</label>
        <input type="month" name="ym" value="{{ ym }}">
        <button type="submit">查看</button>
    </form>

    <hr>

    <h2>📌 本月概况</h2>
    <p>签到义工人数：<b>{{ total_people }}</b></p>
    <p>签到总次数：<b>{{ total_checkins }}</b></p>

    <hr>

    <h2>🏆 义工排行榜</h2>
    <table border="1" cellpadding="8">
        <tr>
            <th>姓名</th>
            <th>签到次数</th>
        </tr>
        {% for r in ranking %}
        <tr>
            <td>{{ r.name }}</td>
            <td>{{ r.checkin_count }}</td>
        </tr>
        {% endfor %}
    </table>

    <hr>

    <h2>🛕 岗位统计</h2>
    <table border="1" cellpadding="8">
        <tr>
            <th>岗位</th>
            <th>次数</th>
        </tr>
        {% for r in place_stats %}
        <tr>
            <td>{{ r.assigned_place or "未安排" }}</td>
            <td>{{ r.total }}</td>
        </tr>
        {% endfor %}
    </table>

    <hr>

    <h2>⚠️ 已排班未签到</h2>
    <table border="1" cellpadding="8">
        <tr>
            <th>日期</th>
            <th>姓名</th>
            <th>角色</th>
            <th>岗位</th>
        </tr>
        {% for r in absent_rows %}
        <tr>
            <td>{{ r.assignment_date }}</td>
            <td>{{ r.name }}</td>
            <td>{{ r.role }}</td>
            <td>{{ r.assigned_place }}</td>
        </tr>
        {% endfor %}
    </table>

    <br>
    <a href="/schedule/admin">返回负责人页面</a>
    """,
    ym=ym,
    total_people=total_people,
    total_checkins=total_checkins,
    ranking=ranking,
    place_stats=place_stats,
    absent_rows=absent_rows
    )


@schedule_bp.route("/schedule/save_day_flags", methods=["POST"])
def save_day_flags():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    date_str = request.form.get("date", "").strip()

    setup_person_1 = request.form.get("setup_person_1", "").strip()
    setup_person_2 = request.form.get("setup_person_2", "").strip()

    remove_person_1 = request.form.get("remove_person_1", "").strip()
    remove_person_2 = request.form.get("remove_person_2", "").strip()
    remove_extra_person = request.form.get("remove_extra_person", "").strip()

    extra_buddha_person = request.form.get("extra_buddha_person", "").strip()

    setup_people = "\n".join(
        x for x in [
            setup_person_1,
            setup_person_2
        ] if x
    )

    remove_people = "\n".join(
        x for x in [
            remove_person_1,
            remove_person_2,
            remove_extra_person
        ] if x
    )

    # 兼容旧字段：有填人员 = 有安排
    need_setup = bool(setup_people)
    need_remove = bool(remove_people)

    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                insert into schedule_day_flags (
                    flag_date,
                    need_setup_master_table,
                    need_remove_master_table,
                    setup_people,
                    remove_people,
                    extra_buddha_person,
                    updated_at
                )
                values (%s, %s, %s, %s, %s, %s, now())
                on conflict (flag_date)
                do update set
                    need_setup_master_table = excluded.need_setup_master_table,
                    need_remove_master_table = excluded.need_remove_master_table,
                    setup_people = excluded.setup_people,
                    remove_people = excluded.remove_people,
                    extra_buddha_person = excluded.extra_buddha_person,
                    updated_at = now()
            """, (
                date_str,
                need_setup,
                need_remove,
                setup_people,
                remove_people,
                extra_buddha_person
            ))

            conn.commit()

    return redirect(url_for(
        "schedule.schedule_admin",
        mode="day",
        override_date=date_str
    ))


@schedule_bp.route("/schedule/fill_pending", methods=["POST"])
def schedule_fill_pending():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    date_str = request.form.get("date", "").strip()

    if not date_str:
        return "❌ 没有日期<br><a href='/schedule/admin'>返回</a>"

    msg = fill_pending_volunteers(date_str)

    return f"""
    <h1>{msg}</h1>
    <a href="/schedule?mode=day&override_date={date_str}">返回当天值班安排</a><br>
    <a href="/schedule/signups?date={date_str}">查看报名管理</a>
    """


def normalize_whatsapp_keyword(text):
    text = str(text or "").strip().upper().replace(" ", "")

    if text.startswith("STW-"):
        return text

    if text.startswith("STW") and text[3:].isdigit():
        return f"STW-{int(text[3:])}"

    if text.startswith("CHE-"):
        return text

    if text.startswith("CHE") and text[3:].isdigit():
        return f"CHE-{int(text[3:])}"

    return text


@schedule_bp.route("/schedule/parse_raw", methods=["POST"])
def schedule_parse_raw():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    single_date = request.form.get("single_date", "").strip()
    raw_signup = request.form.get("raw_signup", "").strip()

    if not single_date:
        return "❌ 没有日期"

    if not raw_signup:
        return "❌ 没有报名内容"

    try:
        signup_date = datetime.strptime(
            single_date,
            "%Y-%m-%d"
        ).date()
    except:
        return "❌ 日期格式错误"

    lines = raw_signup.splitlines()

    added = 0
    skipped = 0

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:

            for line in lines:

                line = line.strip()

                if not line:
                    continue

                parts = line.split()

                if len(parts) < 2:
                    skipped += 1
                    continue

                name = normalize_whatsapp_keyword(parts[0])

                time_text = parts[-1]

                time_text = (
                    time_text
                    .replace("～", "-")
                    .replace("~", "-")
                    .replace("—", "-")
                )

                start_time = "3:00pm"
                end_time = "6:00pm"

                if "-" in time_text:

                    try:
                        s, e = time_text.split("-", 1)

                        start_time = s.strip()

                        if (
                            "am" not in e.lower()
                            and
                            "pm" not in e.lower()
                        ):
                            if "pm" in s.lower():
                                e += "pm"
                            elif "am" in s.lower():
                                e += "am"

                        end_time = e.strip()

                    except:
                        pass

                matches = find_volunteer_by_keyword(name)

                if not matches:
                    skipped += 1
                    continue

                if len(matches) > 1:
                    skipped += 1
                    continue

                vol = matches[0]

                volunteer_id = str(vol["id"])
                real_name = str(vol["name"])

                cur.execute("""
                    select id
                    from volunteer_schedule_signups
                    where volunteer_id = %s
                    and signup_date = %s
                    and role = '值班'
                    and coalesce(status,'pending') <> 'cancelled'
                    limit 1
                """, (
                    volunteer_id,
                    signup_date
                ))

                exists = cur.fetchone()

                if exists:
                    skipped += 1
                    continue

                cur.execute("""
                    insert into volunteer_schedule_signups
                    (
                        volunteer_id,
                        name,
                        signup_date,
                        role,
                        start_time,
                        end_time,
                        status,
                        remarks
                    )
                    values
                    (
                        %s,
                        %s,
                        %s,
                        '值班',
                        %s,
                        %s,
                        'pending',
                        'WhatsApp解析'
                    )
                """, (
                    volunteer_id,
                    real_name,
                    signup_date,
                    start_time,
                    end_time
                ))

                added += 1

        conn.commit()

    return f"""
    <h1>✅ WhatsApp 报名导入完成</h1>

    <p>日期：{single_date}</p>
    <p>成功加入：{added} 位</p>
    <p>跳过：{skipped} 位</p>

    <br>

    <a href="/schedule/admin?mode=day&override_date={single_date}">
        返回当天安排
    </a>
    """


@schedule_bp.route("/schedule/publish", methods=["POST"])
def publish_schedule():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    target_date = request.form["date"]

    publish_schedule_for_date(target_date)

    return redirect(
        url_for(
            "schedule.schedule_admin",
            mode="day",
            override_date=target_date,
            published="1"
        )
    )


@schedule_bp.route("/schedule/override", methods=["POST"])
def schedule_override():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    date = request.form.get("date", "").strip()

    original_names = request.form.getlist("original_name")
    replacement_names = request.form.getlist("replacement_name")

    final_names = []

    for original, replacement in zip(original_names, replacement_names):
        original = str(original).strip()
        replacement = str(replacement).strip()

        if replacement == "__REMOVE__":
            continue

        if replacement:
            final_names.append(replacement)
        elif original:
            final_names.append(original)

    save_buddha_override(date, final_names)

    return redirect(url_for("schedule.schedule_admin", mode="day", override_date=date))


@schedule_bp.route("/schedule/signup_edit/<int:signup_id>", methods=["GET"])
def schedule_signup_edit(signup_id):
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                select *
                from volunteer_schedule_signups
                where id = %s
            """, (signup_id,))
            row = cur.fetchone()

    if not row:
        return "❌ 找不到报名记录<br><a href='/schedule/signups'>返回</a>"

    return render_template_string("""
<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>修改报名</title>
<style>
body { font-family:"Microsoft YaHei", Arial; background:#f5f5f5; padding:20px; }
.box { background:white; max-width:700px; margin:auto; padding:25px; border-radius:15px; }
input, select, button { font-size:22px; padding:8px; margin:8px 0; }
label { display:block; font-size:22px; margin-top:12px; }
</style>
</head>
<body>
<div class="box">

<h1>✏️ 修改报名</h1>

<p><b>{{ row.name }}</b>（{{ row.volunteer_id }}）</p>

<form method="post" action="/schedule/signup_edit/{{ row.id }}">

    <label>日期</label>
    <input type="date" name="signup_date" value="{{ row.signup_date }}" required>

    <label>岗位</label>
    <select name="role" required>
        <option value="值班">值班</option>
        <option value="卫生">卫生</option>
        <option value="设师父供台">设师父供台</option>
        <option value="收师父供台">收师父供台</option>
    </select>

    <label>开始时间</label>
    <select name="start_time">
        {% for t in times %}
        <option value="{{ t }}" {% if row.start_time == t %}selected{% endif %}>{{ t }}</option>
        {% endfor %}
    </select>

    <label>结束时间</label>
    <select name="end_time">
        {% for t in times %}
        <option value="{{ t }}" {% if row.end_time == t %}selected{% endif %}>{{ t }}</option>
        {% endfor %}
    </select>

    <br><br>

    <button type="submit">💾 保存修改</button>
</form>

<br>
<a href="/schedule/signups">⬅ 返回报名管理</a>

</div>
</body>
</html>
""", row=row, roles=ROLE_OPTIONS, times=TIME_OPTIONS)


@schedule_bp.route(
    "/schedule/generate_shortage_notice",
    methods=["GET", "POST"]
)
def generate_shortage_notice():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    date_str = (
        request.form.get("date")
        or request.args.get("date")
    )

    msg = build_shortage_notice_from_assignments(date_str)

    return render_template_string("""
    <!doctype html>
    <html lang="zh">
    <head>
    <meta charset="utf-8">
    <title>缺义工通知</title>
    <meta name="viewport" content="width=device-width, initial-scale=1">

    <link rel="stylesheet"
        href="{{ url_for('static', filename='css/toolbox.css') }}">

    <style>
    .notice-textarea{
        width:100%;
        min-height:420px;
        font-size:20px;
        padding:16px;
        border:1px solid #d1d5db;
        border-radius:16px;
        resize:vertical;
        box-sizing:border-box;
    }
    .notice-actions{
        display:grid;
        grid-template-columns:1fr 1fr;
        gap:14px;
        margin-top:20px;
    }
    @media(max-width:700px){
        .notice-actions{
            grid-template-columns:1fr;
        }
    }
    </style>

    </head>
    <body>

    <div class="page">

        <h1 class="page-title">
            📢 缺义工通知
        </h1>

        <p class="page-subtitle">
            WhatsApp 缺义工通知，可直接复制发送
        </p>

        <div class="card">

            <div class="section-title">
                📄 通知内容
            </div>

            <textarea
                id="notice_text"
                class="notice-textarea"
            >{{ msg }}</textarea>

            <div class="notice-actions">

                <button
                    class="btn-tool btn-primary"
                    onclick="copyNotice()">

                    📋 复制通知

                </button>

                <a
                    class="btn-tool btn-secondary"
                    href="/schedule?mode=day&override_date={{ date_str }}">

                    ← 返回负责人首页

                </a>

            </div>

        </div>

    </div>

    <script>
    function copyNotice(){

        const ta = document.getElementById("notice_text");

        if(!ta){
            alert("❌ 找不到通知内容");
            return;
        }

        ta.focus();
        ta.select();
        ta.setSelectionRange(0, 999999);

        const ok = document.execCommand("copy");

        if(ok){
            alert("✅ 已复制通知");
        }else{
            alert("❌ 复制失败，请手动全选复制");
        }
    }
    </script>

    </body>
    </html>
    """,
    msg=msg,
    date_str=date_str
    )


@schedule_bp.route("/schedule/view_assigned", methods=["POST"])
def schedule_view_assigned():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    date_str = request.form.get("date", "").strip()

    rows = load_assigned_places_for_date(date_str)

    cleaning_groups = {
        "佛堂卫生": [],
        "二楼卫生": [],
        "楼梯卫生": [],
    }

    offering_group = []

    duty_groups = {
        "绿": {
            "观音堂": [],
            "活动中心": [],
        },
        "橙": {
            "观音堂": [],
            "活动中心": [],
        },
        "黄": {
            "观音堂": [],
            "活动中心": [],
        },
    }

    other_groups = {}

    for r in rows:
        item = {
            "id": r["id"],
            "name": r["name"],
            "role": r.get("role"),
            "shift_label": r.get("shift_label"),
            "assigned_place": r.get("assigned_place"),
            "start_time": r.get("start_time"),
            "end_time": r.get("end_time"),
        }

        place = r.get("assigned_place") or "未安排"
        role = r.get("role")
        shift = r.get("shift_label")

        if role == "卫生" and place in cleaning_groups:
            cleaning_groups[place].append(item)

        elif role == "供台" or place == "设师父供台":
            offering_group.append(item)

        elif role == "值班":
            shift_key = str(shift or "").replace("班", "")

            if shift_key in duty_groups and place in duty_groups[shift_key]:
                duty_groups[shift_key][place].append(item)
            else:
                other_groups.setdefault(place, []).append(item)

        else:
            other_groups.setdefault(place, []).append(item)

    return render_template_string("""
    <!doctype html>
    <html>
    <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>排班结果</title>

    <link rel="stylesheet"
          href="{{ url_for('static', filename='css/toolbox.css') }}">

    <style>
    .assigned-wrap{
        max-width:1100px;
        margin:auto;
        padding:24px;
    }

    .assignment-section{
        background:#f8fafc;
        border:1px solid #e5e7eb;
        border-radius:18px;
        padding:20px;
        margin:16px 0;
    }

    .assignment-place-title{
        font-size:30px;
        font-weight:bold;
        margin-bottom:14px;
    }

    .person-row{
        display:flex;
        justify-content:space-between;
        align-items:center;
        gap:16px;
        border-bottom:1px solid #e5e7eb;
        padding:14px 0;
        font-size:28px;
    }

    .person-row:last-child{
        border-bottom:none;
    }

    .person-name{
        font-weight:bold;
    }

    .person-time{
        color:#6b7280;
        font-size:24px;
    }

    .edit-link{
        font-size:22px;
        padding:8px 14px;
        border-radius:12px;
        background:#dbeafe;
        color:#1e40af;
        font-weight:bold;
        white-space:nowrap;
    }

    .shift-title{
        font-size:34px;
        margin-top:26px;
    }

    .empty-line{
        font-size:24px;
        color:#6b7280;
    }
    </style>
    </head>

    <body>

    <div class="assigned-wrap">

        <div class="card">

            <h1 class="page-title">
                📅 {{ date_str }} 排班结果
            </h1>

            <p class="page-subtitle">
                查看正式安排，也可进入个别义工调整岗位。
            </p>

            <div class="btn-row">

                <a class="btn-tool btn-gray"
                href="/schedule?mode=day&override_date={{ date_str }}">
                    ⬅ 返回当天值班安排
                </a>

                <form method="post" action="/schedule/copy_whatsapp">
                    <input type="hidden" name="date" value="{{ date_str }}">
                    <button class="btn-tool btn-green" type="submit">
                        📋 生成 WhatsApp 文字
                    </button>
                </form>

            </div>

        </div>

        <div class="card">

            <h2 class="section-title">🧹 卫生</h2>

            {% for place, names in cleaning_groups.items() %}
                <div class="assignment-section">

                    <div class="assignment-place-title">
                        {{ place }}
                    </div>

                    {% if names %}
                        {% for item in names %}
                            <div class="person-row">
                                <div>
                                    <span class="person-name">{{ item.name }}</span>
                                </div>

                                <a class="edit-link"
                                href="/schedule/edit_assigned/{{ item.id }}">
                                    ✏️ 调整
                                </a>
                            </div>
                        {% endfor %}
                    {% else %}
                        <div class="empty-line">暂无安排</div>
                    {% endif %}

                </div>
            {% endfor %}

        </div>

        <div class="card">

            <h2 class="section-title">🙏 供台</h2>

            <div class="assignment-section">

                {% if offering_group %}
                    {% for item in offering_group %}
                        <div class="person-row">
                            <div>
                                <span class="person-name">{{ item.name }}</span>
                            </div>

                            <a class="edit-link"
                            href="/schedule/edit_assigned/{{ item.id }}">
                                ✏️ 调整
                            </a>
                        </div>
                    {% endfor %}
                {% else %}
                    <div class="empty-line">暂无安排</div>
                {% endif %}

            </div>

        </div>

        <div class="card">

            <h2 class="section-title">🏠 值班</h2>

            {% for shift, places in duty_groups.items() %}

                {% if shift == "绿" %}
                    <h3 class="shift-title">🟢 绿班</h3>
                {% elif shift == "橙" %}
                    <h3 class="shift-title">🟠 橙班</h3>
                {% elif shift == "黄" %}
                    <h3 class="shift-title">🟡 黄班</h3>
                {% else %}
                    <h3 class="shift-title">{{ shift }}</h3>
                {% endif %}

                {% for place, names in places.items() %}
                    <div class="assignment-section">

                        <div class="assignment-place-title">
                            {{ place }}
                        </div>

                        {% if names %}
                            {% for item in names %}
                                <div class="person-row">

                                    <div>
                                        <span class="person-name">
                                            {{ item.name }}
                                        </span>

                                        {% if item.start_time and item.end_time %}
                                            <span class="person-time">
                                                （{{ item.start_time }} ~ {{ item.end_time }}）
                                            </span>
                                        {% endif %}
                                    </div>

                                    <a class="edit-link"
                                    href="/schedule/edit_assigned/{{ item.id }}">
                                        ✏️ 调整
                                    </a>

                                </div>
                            {% endfor %}
                        {% else %}
                            <div class="empty-line">暂无安排</div>
                        {% endif %}

                    </div>
                {% endfor %}

            {% endfor %}

        </div>

        {% if other_groups %}
        <div class="card">

            <h2 class="section-title">其他</h2>

            {% for place, names in other_groups.items() %}
                <div class="assignment-section">

                    <div class="assignment-place-title">
                        {{ place }}
                    </div>

                    {% for item in names %}
                        <div class="person-row">

                            <div>
                                <span class="person-name">
                                    {{ item.name }}
                                </span>

                                {% if item.start_time and item.end_time %}
                                    <span class="person-time">
                                        （{{ item.start_time }} ~ {{ item.end_time }}）
                                    </span>
                                {% endif %}
                            </div>

                            <a class="edit-link"
                            href="/schedule/edit_assigned/{{ item.id }}">
                                ✏️ 调整
                            </a>

                        </div>
                    {% endfor %}

                </div>
            {% endfor %}

        </div>
        {% endif %}

    </div>

    </body>
    </html>
    """,
        date_str=date_str,
        cleaning_groups=cleaning_groups,
        offering_group=offering_group,
        duty_groups=duty_groups,
        other_groups=other_groups,
    )


@schedule_bp.route("/reports")
def public_reports():

    supabase = get_supabase_client()

    files = supabase.storage.from_("reports").list()

    report_files = []

    for f in files:
        filename = f.get("name")

        if not filename or not filename.endswith(".xlsx"):
            continue

        signed = supabase.storage.from_("reports").create_signed_url(
            filename,
            60 * 60
        )

        report_files.append({
            "filename": filename,
            "display_name": display_report_name(filename),
            "url": signed["signedURL"]
        })

    report_files.sort(
        key=lambda x: x["filename"],
        reverse=True
    )

    return render_template_string("""
    <!doctype html>
    <html>
    <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>文件中心</title>

    <link rel="stylesheet"
          href="{{ url_for('static', filename='css/toolbox.css') }}">

    <style>
    .reports-wrap{
        max-width:900px;
        margin:auto;
        padding:24px;
    }

    .file-row{
        display:flex;
        justify-content:space-between;
        align-items:center;
        gap:18px;
        padding:18px 0;
        border-bottom:1px solid #e5e7eb;
        font-size:26px;
    }

    .file-row:last-child{
        border-bottom:none;
    }

    .file-name{
        font-weight:bold;
    }

    .download-btn{
        min-width:150px;
    }

    @media (max-width:700px){
        .file-row{
            flex-direction:column;
            align-items:stretch;
        }

        .download-btn{
            width:100%;
        }
    }
    </style>
    </head>

    <body>

    <div class="reports-wrap">

        <div class="card">

            <h1 class="page-title">📚 观音堂资料中心</h1>

            <p class="page-subtitle">
                提供月报、年报及相关资料下载。
            </p>

            {% if report_files %}

                {% for f in report_files %}
                <div class="file-row">

                    <div class="file-name">
                        📄 {{ f.display_name }}
                    </div>

                    <a class="btn-tool btn-blue download-btn"
                    href="{{ f.url }}">
                        📥 下载
                    </a>

                </div>
                {% endfor %}

            {% else %}

                <div class="empty-state">
                    <div class="empty-icon">📭</div>
                    <div class="empty-title">
                        目前还没有上传报表
                    </div>
                    <div class="empty-text">
                        上传月报或年报后，会显示在这里。
                    </div>
                </div>

            {% endif %}

            <div class="btn-row">
                <a class="btn-tool btn-gray btn-full"
                href="/schedule/admin">
                    ⬅ 返回负责人页面
                </a>
            </div>

        </div>

    </div>

    </body>
    </html>
    """, report_files=report_files)


@schedule_bp.route("/schedule/notice_center")
def schedule_notice_center():
    
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    now = malaysia_now()

    if now.hour >= 18:
        default_date = (now + timedelta(days=1)).strftime("%Y-%m-%d")
    else:
        default_date = now.strftime("%Y-%m-%d")

    date_str = request.args.get("date") or default_date
    republish_info = get_schedule_republish_info(date_str)

    return render_template_string("""
<!doctype html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>
body{
    font-family:"Microsoft YaHei";
    background:#f5f5f5;
    padding:20px;
    font-size:26px;
}
.box{
    max-width:900px;
    margin:auto;
    background:white;
    border-radius:20px;
    padding:30px;
}
.notice-btn{
    display:flex;
    flex-direction:column;
    justify-content:center;
    align-items:center;

    text-align:center;

    font-size:34px;
    font-weight:bold;

    height:110px;
}

.notice-btn small{
    margin-top:8px;
    font-size:20px;
    font-weight:normal;
    opacity:0.9;
}
.notice-btn{
    display:block;
    width:100%;
    text-align:center;
    text-decoration:none;
    color:white;
    border-radius:18px;
    padding:28px;
    margin:18px 0;
    font-size:32px;
    font-weight:bold;
    box-sizing:border-box;
}
.blue{background:#2196F3;}
.green{background:#25D366;}
.orange{background:#FF9800;}
.gray{background:#607D8B;}
</style>
</head>
<body>
<div class="box">

<h1>📢 通知中心</h1>

<p style="font-size:22px;color:#666;">
负责人每天只需依照以下顺序发送 WhatsApp 信息。
</p>
                                  
{% if republish_info is defined and republish_info.need_republish %}
<div style="
background:#fff8dc;
border:3px solid orange;
border-radius:15px;
padding:18px;
margin-bottom:20px;
font-size:24px;
line-height:1.8;
">

🟠 <b>正式值班表已有更新</b><br>

自上次发布正式值班表后，已有义工报名、修改报名、取消报名或负责人调整排班。

<br><br>

<b>请重新发送最新正式值班表 WhatsApp。</b>

<br><br>

<form method="post" action="/schedule/copy_whatsapp">
    <input type="hidden" name="date" value="{{ date_str }}">

    <button class="notice-btn green" type="submit">
        📋 复制最新正式值班表 WhatsApp<br>
        <small>📤 可随时重新发送最新版</small>
    </button>
</form>

{% if republish_info.last_schedule_change_text %}
<br>
<small>
    🕒 最后更新：{{ republish_info.last_schedule_change_text }}
</small>
{% endif %}

</div>
{% endif %}

<a class="notice-btn blue"
   href="/schedule/signup_notice?date={{ date_str }}">

    📢 报名通知<br>
    <small>🕖 晚上 7:00 发群</small>

</a>

<a class="notice-btn orange"
   href="/schedule/generate_shortage_notice?date={{ date_str }}">

    📣 缺义工通知<br>
    <small>📌 缺人时发送</small>

</a>

<a class="notice-btn gray" href="/schedule?mode=day&override_date={{ date_str }}">
   ⬅ 返回负责人首页
</a>

</div>
</body>
</html>
""", date_str=date_str,republish_info=republish_info)


@schedule_bp.route("/schedule/signup_notice")
def schedule_signup_notice():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    now = malaysia_now()

    if now.hour >= 18:
        default_date = (now + timedelta(days=1)).strftime("%Y-%m-%d")
    else:
        default_date = now.strftime("%Y-%m-%d")

    date_str = default_date

    text = build_signup_notice_whatsapp(date_str)

    return render_template_string("""
    <!doctype html>
    <html lang="zh">
    <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>今晚义工报名</title>

    <link rel="stylesheet"
        href="{{ url_for('static', filename='css/toolbox.css') }}">

    <style>

    .notice-textarea{
        width:100%;
        min-height:520px;
        font-size:19px;
        line-height:1.6;
        padding:18px;
        border:1px solid #d1d5db;
        border-radius:16px;
        box-sizing:border-box;
        resize:vertical;
    }

    .notice-actions{
        display:grid;
        grid-template-columns:1fr 1fr;
        gap:14px;
        margin-top:20px;
    }

    @media(max-width:700px){

        .notice-actions{
            grid-template-columns:1fr;
        }

    }

    </style>

    </head>
    <body>

    <div class="page">

        <h1 class="page-title">
            📢 今晚义工报名
        </h1>

        <p class="page-subtitle">
            WhatsApp 报名通知，可直接复制发送
        </p>

        <div class="card">

            <div class="section-title">
                📄 通知内容
            </div>

            <textarea
                id="noticeText"
                class="notice-textarea"
                readonly>{{ text }}</textarea>

            <div class="notice-actions">

                <button
                    type="button"
                    class="btn-tool btn-primary"
                    onclick="copyText()">

                    📋 复制 WhatsApp

                </button>

                <a
                    class="btn-tool btn-secondary"
                    href="/schedule/notice_center?date={{ date_str }}">

                    ← 返回通知中心

                </a>

            </div>

        </div>

    </div>

    <script>

    function copyText(){

        const ta = document.getElementById("noticeText");

        ta.focus();
        ta.select();
        ta.setSelectionRange(0,999999);

        const ok = document.execCommand("copy");

        if(ok){
            alert("✅ 已复制，可以贴去 WhatsApp 群");
        }else{
            alert("❌ 复制失败，请手动复制");
        }

    }

    </script>

    </body>
    </html>

    """, text=text, date_str=date_str)


@schedule_bp.route("/schedule/signup_edit/<int:signup_id>", methods=["POST"])
def schedule_signup_edit_save(signup_id):
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    signup_date = request.form.get("signup_date", "").strip()
    role = request.form.get("role", "").strip()
    start_time = request.form.get("start_time", "").strip()
    end_time = request.form.get("end_time", "").strip()

    if role == "值班":
        if not start_time or not end_time:
            return "❌ 值班必须选择开始和结束时间<br><a href='/volunteer/signup'>返回</a>"

        if end_time <= start_time:
            return "❌ 结束时间必须比开始时间迟<br><a href='/volunteer/signup'>返回</a>"
    else:
        start_time = None
        end_time = None

    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                update volunteer_schedule_signups
                set
                    signup_date = %s,
                    role = %s,
                    start_time = %s,
                    end_time = %s,
                    status = 'pending',
                    assigned_place = null,
                    remarks = '负责人修改报名'
                where id = %s
            """, (
                signup_date,
                role,
                start_time,
                end_time,
                signup_id
            ))

            conn.commit()

    sync_schedule_after_signup_change(
        signup_id,
        action="upsert",
        changed_by="admin"
    )

    return redirect("/schedule/signups")

@schedule_bp.route("/schedule/signup_place/<int:signup_id>", methods=["GET"])
def schedule_signup_place(signup_id):
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:

            cur.execute("""
                select *
                from volunteer_schedule_signups
                where id = %s
            """, (signup_id,))
            signup = cur.fetchone()

            if not signup:
                return "❌ 找不到报名记录<br><a href='/schedule/signups'>返回</a>"

            cur.execute("""
                select id
                from volunteer_schedule_assignments
                where signup_id = %s
                and coalesce(status, 'assigned') <> 'cancelled'
                order by start_time, id
            """, (signup_id,))
            assignments = cur.fetchall()

    if len(assignments) == 1:
        assignment_id = assignments[0]["id"]
        return redirect(f"/schedule/edit_assigned/{assignment_id}")

    if len(assignments) > 1:
        return f"""
        <h1>⚠️ 这位义工有多段安排</h1>
        <p>{signup['name']} 在 {signup['signup_date']} 有多段排班。</p>
        <p>请到当天排班结果页面逐段调整。</p>

        <form method="post" action="/schedule/view_assigned">
            <input type="hidden" name="date" value="{signup['signup_date']}">
            <button type="submit">查看当天排班结果</button>
        </form>

        <br>
        <a href="/schedule/signups">返回报名管理</a>
        """

    return f"""
    <h1>❌ 还没有真正排班记录</h1>
    <p>{signup['name']} 目前在 assignments 里没有安排。</p>
    <p>请先自动排班，或到排班结果页面补入。</p>
    <a href="/schedule/signups">返回报名管理</a>
    """


@schedule_bp.route("/schedule/copy_month_prebook", methods=["POST"])
def copy_month_prebook():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    year = int(request.form.get("year"))
    month = int(request.form.get("month"))

    msg = build_month_prebook_message(year, month)

    return render_template_string("""
    <h1>📋 {{ year }}年{{ month }}月预报名名单</h1>

    <textarea
        id="msg"
        style="width:100%;height:600px;font-size:22px;"
    >{{ msg }}</textarea>

    <br><br>

    <button
        onclick="navigator.clipboard.writeText(document.getElementById('msg').value)"
        style="font-size:24px;padding:15px 25px;"
    >
        📋 复制
    </button>

    <br><br>

    <a href="/schedule?mode=prebook">返回月预报名</a>
    """, msg=msg, year=year, month=month)


@schedule_bp.route("/schedule/attendance_status")
def schedule_attendance_status():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    today = malaysia_today()

    date_str = request.args.get(
        "date",
        today.strftime("%Y-%m-%d")
    )

    def group_by_person(rows, source="assignment"):
        grouped = {}

        for r in rows:
            key = r.get("volunteer_id") or r.get("name")

            if not key:
                continue

            if key not in grouped:
                grouped[key] = {
                    "name": r.get("name"),
                    "volunteer_id": r.get("volunteer_id"),
                    "roles": [],
                    "places": [],
                }

            role = r.get("role") or r.get("actual_role") or ""
            if role and role not in grouped[key]["roles"]:
                grouped[key]["roles"].append(role)

            if source == "log":
                place = r.get("actual_place") or ""
                shift = ""
            else:
                place = r.get("assigned_place") or ""
                shift = r.get("shift_label") or ""

            place_text = f"{shift} {place}".strip()

            if place_text and place_text not in grouped[key]["places"]:
                grouped[key]["places"].append(place_text)

        result = []

        for g in grouped.values():
            result.append({
                "name": g["name"],
                "volunteer_id": g["volunteer_id"],
                "role_text": " / ".join(g["roles"]),
                "place_text": " / ".join(g["places"]),
            })

        return result

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:

            cur.execute("""
                select *
                from volunteer_schedule_assignments
                where assignment_date = %s
                and coalesce(status, 'assigned') <> 'cancelled'
                order by role, shift_label, assigned_place, start_time, name
            """, (date_str,))
            assignments = cur.fetchall()

            cur.execute("""
                select *
                from volunteer_attendance_logs
                where attendance_date = %s
                order by check_in_time, name
            """, (date_str,))
            logs = cur.fetchall()

    signed_assignment_ids = {
        r["assignment_id"]
        for r in logs
        if r.get("assignment_id")
    }

    checked_in_raw = []
    not_checked_in_raw = []
    walk_ins_raw = []
    not_checked_out_raw = []

    for a in assignments:
        if a["id"] in signed_assignment_ids:
            checked_in_raw.append(a)
        else:
            not_checked_in_raw.append(a)

    for l in logs:
        if l.get("walk_in"):
            walk_ins_raw.append(l)

        if not l.get("check_out_time"):
            not_checked_out_raw.append(l)

    checked_in = group_by_person(checked_in_raw, source="assignment")
    not_checked_in = group_by_person(not_checked_in_raw, source="assignment")
    walk_ins = group_by_person(walk_ins_raw, source="log")
    not_checked_out = group_by_person(not_checked_out_raw, source="log")

    return render_template_string("""
    <!doctype html>
    <html>
    <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>签到情况</title>

    <link rel="stylesheet"
          href="{{ url_for('static', filename='css/toolbox.css') }}">

    <style>
    .attendance-wrap{
        max-width:1100px;
        margin:auto;
        padding:24px;
    }

    .attendance-list{
        list-style:none;
        padding:0;
        margin:0;
    }

    .attendance-list li{
        font-size:26px;
        padding:14px 0;
        border-bottom:1px solid #e5e7eb;
    }

    .attendance-list li:last-child{
        border-bottom:none;
    }

    .date-query-row{
        display:grid;
        grid-template-columns:260px 180px;
        gap:14px;
        align-items:end;
    }
                                  
    .date-query-row .form-group{
        margin-bottom:0;
    }

    .date-query-row .btn-tool{
        height:68px;
        min-height:68px;
        font-size:26px;
        align-self:end;
    }

    @media (max-width:700px){
        .date-query-row{
            grid-template-columns:1fr;
        }
    }
    </style>
    </head>

    <body>

    <div class="attendance-wrap">

        <div class="card">

            <h1 class="page-title">
                📋 签到情况
            </h1>

            <p class="page-subtitle">
                日期：{{ date_str }}
            </p>

            <form method="get">

                <div class="date-query-row">

                    <div class="form-group">
                        <label class="form-label">查询日期</label>

                        <input
                            class="form-input"
                            type="date"
                            name="date"
                            value="{{ date_str }}">
                    </div>

                    <button
                        class="btn-tool btn-blue"
                        type="submit">
                          查询
                    </button>

                </div>

            </form>

            <div class="btn-row">
                <a class="btn-tool btn-gray"
                href="/schedule/admin">
                    ⬅ 返回负责人首页
                </a>
            </div>

        </div>

        <div class="summary-grid">

            <div class="summary-box">
                <div class="summary-title">✅ 已签到</div>
                <div class="summary-value">{{ checked_in|length }}</div>
            </div>

            <div class="summary-box">
                <div class="summary-title">❌ 未签到</div>
                <div class="summary-value">{{ not_checked_in|length }}</div>
            </div>

            <div class="summary-box">
                <div class="summary-title">🚶 临时报到</div>
                <div class="summary-value">{{ walk_ins|length }}</div>
            </div>

            <div class="summary-box">
                <div class="summary-title">⏳ 未签退</div>
                <div class="summary-value">{{ not_checked_out|length }}</div>
            </div>

        </div>

        <div class="card">
            <h2 class="section-title">✅ 已签到：{{ checked_in|length }}</h2>

            <ul class="attendance-list">
            {% for r in checked_in %}
                <li>{{ r.name }}｜{{ r.role_text }}｜{{ r.place_text }}</li>
            {% else %}
                <li class="text-gray">暂无记录</li>
            {% endfor %}
            </ul>
        </div>

        <div class="card">
            <h2 class="section-title">❌ 已排班未签到：{{ not_checked_in|length }}</h2>

            <ul class="attendance-list">
            {% for r in not_checked_in %}
                <li>{{ r.name }}｜{{ r.role_text }}｜{{ r.place_text }}</li>
            {% else %}
                <li class="text-gray">暂无记录</li>
            {% endfor %}
            </ul>
        </div>

        <div class="card">
            <h2 class="section-title">🚶 临时报到：{{ walk_ins|length }}</h2>

            <ul class="attendance-list">
            {% for r in walk_ins %}
                <li>{{ r.name }}｜{{ r.role_text }}｜{{ r.place_text }}</li>
            {% else %}
                <li class="text-gray">暂无记录</li>
            {% endfor %}
            </ul>
        </div>

        <div class="card">
            <h2 class="section-title">⏳ 未签退：{{ not_checked_out|length }}</h2>

            <ul class="attendance-list">
            {% for r in not_checked_out %}
                <li>{{ r.name }}｜{{ r.role_text }}｜{{ r.place_text }}</li>
            {% else %}
                <li class="text-gray">暂无记录</li>
            {% endfor %}
            </ul>
        </div>

    </div>

    </body>
    </html>
    """,
        date_str=date_str,
        checked_in=checked_in,
        not_checked_in=not_checked_in,
        walk_ins=walk_ins,
        not_checked_out=not_checked_out,
    )


@schedule_bp.route("/schedule/settings", methods=["GET", "POST"])
def schedule_settings():

    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    if request.method == "POST":
        keys = [
            "default_day_switch_time",
            "target_orange_guanyintang",
            "target_orange_activity",
            "target_yellow_guanyintang",
            "target_yellow_activity",
            "target_cleaning",
            "supply_alert_days",
        ]

        for key in keys:
            value = request.form.get(key, "").strip()
            save_schedule_setting(key, value)

        return redirect(url_for("schedule.schedule_settings", saved="1"))

    settings = get_schedule_settings()
    saved = request.args.get("saved") == "1"
    daily_quote = get_daily_buddha_quote()

    return render_template_string("""
    <!doctype html>
    <html>
    <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>排班设置</title>

    <link rel="stylesheet"
          href="{{ url_for('static', filename='css/toolbox.css') }}">

    <style>
    .settings-wrap{
        max-width:900px;
        margin:auto;
        padding:24px;
    }

    .settings-grid{
        display:grid;
        grid-template-columns:1fr 1fr;
        gap:18px;
    }

    @media(max-width:700px){
        .settings-grid{
            grid-template-columns:1fr;
        }
    }
    </style>
    </head>

    <body>

    <div class="settings-wrap">

        <div class="card">

            <h1 class="page-title">⚙️ 排班系统设置</h1>

            <p class="page-subtitle">
                管理负责人页面默认日期、缺义工目标人数与供台提醒。
            </p>

            <div class="quote-card" style="margin-bottom:22px;">
                <div class="quote-title">
                    🌸 今日佛言佛语
                </div>

                <div class="quote-content">
                    {{ daily_quote }}
                </div>
            </div>

            {% if saved %}
            <div class="alert alert-success">
                ✅ 设置已保存
            </div>
            {% endif %}

            <form method="post">

                <h2 class="section-title">📅 默认工作日期</h2>

                <div class="form-group">
                    <label class="form-label">
                        几点后默认看明天？
                    </label>

                    <input
                        class="form-input"
                        name="default_day_switch_time"
                        value="{{ settings.default_day_switch_time }}"
                        placeholder="18:00">

                    <div class="form-help">
                        例：18:00 = 晚上 6 点后，负责人首页默认切换到明天。
                    </div>
                </div>

                <div class="divider"></div>

                <h2 class="section-title">👥 缺义工目标人数</h2>

                <div class="settings-grid">

                    <div class="form-group">
                        <label class="form-label">橙班 - 观音堂</label>
                        <input
                            class="form-input"
                            name="target_orange_guanyintang"
                            value="{{ settings.target_orange_guanyintang }}">
                    </div>

                    <div class="form-group">
                        <label class="form-label">橙班 - 活动中心</label>
                        <input
                            class="form-input"
                            name="target_orange_activity"
                            value="{{ settings.target_orange_activity }}">
                    </div>

                    <div class="form-group">
                        <label class="form-label">黄班 - 观音堂</label>
                        <input
                            class="form-input"
                            name="target_yellow_guanyintang"
                            value="{{ settings.target_yellow_guanyintang }}">
                    </div>

                    <div class="form-group">
                        <label class="form-label">黄班 - 活动中心</label>
                        <input
                            class="form-input"
                            name="target_yellow_activity"
                            value="{{ settings.target_yellow_activity }}">
                    </div>

                    <div class="form-group">
                        <label class="form-label">卫生人数</label>
                        <input
                            class="form-input"
                            name="target_cleaning"
                            value="{{ settings.target_cleaning }}">
                    </div>

                </div>

                <div class="divider"></div>

                <h2 class="section-title">🪔 供台提醒</h2>

                <div class="form-group">
                    <label class="form-label">
                        提醒未来几天的大日子？
                    </label>

                    <input
                        class="form-input"
                        name="supply_alert_days"
                        value="{{ settings.supply_alert_days }}">

                    <div class="form-help">
                        系统会提前显示未来几天需要设供台的大日子。
                    </div>
                </div>

                <div class="btn-row">

                    <button
                        class="btn-tool btn-green btn-full"
                        type="submit">
                        保存排班设置
                    </button>

                    <a
                        class="btn-tool btn-gray btn-full"
                        href="{{ url_for('schedule.schedule_admin') }}">
                        返回负责人首页
                    </a>

                </div>

            </form>

        </div>

    </div>

    </body>
    </html>
    """, settings=settings, saved=saved, daily_quote=daily_quote)


@schedule_bp.route("/schedule/signups")
def schedule_signups_manage():
    if not session.get("schedule_login"):
        return redirect(url_for("schedule.schedule_admin"))

    status = request.args.get("status", "all")
    date_filter = request.args.get("date", "").strip()
    role_filter = request.args.get("role", "all")
    show_history = request.args.get("show_history") == "1"

    sql = """
        select
            id,
            volunteer_id,
            name,
            signup_date,
            role,
            start_time,
            end_time,
            coalesce(status, 'pending') as status,
            assigned_place,
            remarks
        from volunteer_schedule_signups
        where 1=1
    """

    if not show_history:
        sql += """
            and signup_date >= current_date
            and coalesce(status,'pending') <> 'cancelled'
        """

    params = []

    if status != "all":
        sql += " and coalesce(status, 'pending') = %s"
        params.append(status)

    if date_filter:
        sql += " and signup_date = %s"
        params.append(date_filter)

    if role_filter != "all":
        sql += " and role = %s"
        params.append(role_filter)

    sql += " order by signup_date, role, start_time, name"

    with get_conn() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(sql, params)
            rows = cur.fetchall()

            signup_ids = [r["id"] for r in rows]
            assignment_map = {}

            if signup_ids:
                cur.execute("""
                    select
                        signup_id,
                        shift_label,
                        assigned_place,
                        start_time,
                        end_time
                    from volunteer_schedule_assignments
                    where signup_id = any(%s)
                    and coalesce(status, 'assigned') <> 'cancelled'
                    order by start_time, shift_label, assigned_place
                """, (signup_ids,))

                assignment_rows = cur.fetchall()

                for a in assignment_rows:
                    sid = a["signup_id"]

                    shift = a.get("shift_label") or ""
                    place = a.get("assigned_place") or ""
                    start = a.get("start_time") or ""
                    end = a.get("end_time") or ""

                    text = ""

                    if shift:
                        text += shift

                    if place:
                        if text:
                            text += " "
                        text += place

                    if start and end:
                        text += f"<br><small>{start} ~ {end}</small>"

                    assignment_map.setdefault(sid, []).append(text)

            for r in rows:
                r["final_assignment"] = assignment_map.get(r["id"], [])

    summary = {
        "total": len(rows),
        "pending": sum(1 for r in rows if r["status"] == "pending"),
        "assigned": sum(1 for r in rows if r["status"] == "assigned"),
        "cancelled": sum(1 for r in rows if r["status"] == "cancelled"),
    }

    today_date = malaysia_today()

    today = today_date.strftime("%Y-%m-%d")
    tomorrow = (today_date + timedelta(days=1)).strftime("%Y-%m-%d")

    role_summary = {}

    for r in rows:
        role = r["role"] or "未填写"
        role_summary[role] = role_summary.get(role, 0) + 1

    return render_template_string("""
    <!doctype html>
    <html>
    <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>报名管理</title>

    <link rel="stylesheet"
          href="{{ url_for('static', filename='css/toolbox.css') }}">

    <style>
    .signup-admin-wrap{
        max-width:1280px;
        margin:auto;
        padding:24px;
    }

    .filter-grid{
        display:grid;
        grid-template-columns:240px 1fr 1fr 180px;
        gap:16px;
        align-items:end;
    }

    .quick-date-row{
        display:flex;
        flex-wrap:wrap;
        gap:12px;
        margin-top:16px;
        align-items:center;
    }

    .role-stat-row{
        display:flex;
        flex-wrap:wrap;
        gap:12px;
        margin-top:12px;
    }

    .role-stat-pill{
        background:#fff7ed;
        color:#9a3412;
        border:1px solid #fed7aa;
        border-radius:999px;
        padding:10px 18px;
        font-size:22px;
        font-weight:bold;
    }

    .status-pending{
        background:#fff8d6;
    }

    .status-assigned{
        background:#e8fff0;
    }

    .status-cancelled{
        background:#f3f4f6;
        color:#777;
    }

    .signup-table{
        width:100%;
        border-collapse:collapse;
        font-size:22px;
        background:white;
    }

    .signup-table th{
        background:#f8fafc;
        font-size:22px;
        padding:14px 10px;
        border-bottom:2px solid #e5e7eb;
        white-space:nowrap;
    }

    .signup-table td{
        padding:14px 10px;
        border-bottom:1px solid #e5e7eb;
        text-align:center;
        vertical-align:top;
    }

    .signup-table tr:hover{
        filter:brightness(.98);
    }

    .assignment-text{
        font-size:20px;
        line-height:1.6;
    }

    .remark-text{
        font-size:20px;
        color:#6b7280;
        max-width:200px;
    }

    .op-stack{
        display:flex;
        flex-direction:column;
        gap:8px;
        align-items:center;
    }

    .op-stack form{
        margin:0;
    }

    .op-btn{
        min-height:42px;
        padding:0 14px;
        font-size:18px;
        border-radius:12px;
        width:110px;
    }

    .floating-back{
        position:fixed;
        right:25px;
        bottom:25px;
        z-index:9999;
        min-width:170px;
    }

    .history-check{
        display:flex;
        gap:8px;
        align-items:center;
        font-size:22px;
        font-weight:bold;
        color:#374151;
    }

    .history-check input{
        width:22px;
        height:22px;
    }

    @media (max-width:900px){
        .filter-grid{
            grid-template-columns:1fr;
        }

        .signup-table{
            font-size:20px;
        }

        .signup-table th,
        .signup-table td{
            padding:10px 8px;
        }

        .floating-back{
            position:static;
            margin:20px auto;
            width:100%;
        }
    }
    </style>
    </head>

    <body>

    <div class="signup-admin-wrap">

        <div class="card">

            <h1 class="page-title">📋 义工报名管理</h1>

            <p class="page-subtitle">
                查看、筛选、修改、取消或恢复义工报名。
            </p>

            <div class="btn-row">
                <a class="btn-tool btn-gray"
                href="/schedule/admin">
                    ⬅ 返回负责人首页
                </a>
            </div>

        </div>

        <div class="summary-grid">

            <div class="summary-box">
                <div class="summary-title">当前查询</div>
                <div class="summary-value">{{ summary.total }}</div>
            </div>

            <div class="summary-box">
                <div class="summary-title">🟡 未安排</div>
                <div class="summary-value">{{ summary.pending }}</div>
            </div>

            <div class="summary-box">
                <div class="summary-title">🟢 已安排</div>
                <div class="summary-value">{{ summary.assigned }}</div>
            </div>

            <div class="summary-box">
                <div class="summary-title">⚪ 已取消</div>
                <div class="summary-value">{{ summary.cancelled }}</div>
            </div>

        </div>

        <div class="card">

            <h2 class="section-title">📊 岗位统计</h2>

            {% if role_summary %}
                <div class="role-stat-row">
                {% for role, count in role_summary.items() %}
                    <div class="role-stat-pill">
                        {{ role }}：{{ count }} 人
                    </div>
                {% endfor %}
                </div>
            {% else %}
                <div class="empty-state">
                    <div class="empty-icon">📭</div>
                    <div class="empty-title">没有岗位统计</div>
                </div>
            {% endif %}

        </div>

        <div class="card">

            <h2 class="section-title">🔍 筛选查询</h2>

            <form method="get">

                <div class="filter-grid">

                    <div class="form-group">
                        <label class="form-label">日期</label>
                        <input
                            class="form-input"
                            type="date"
                            name="date"
                            value="{{ date_filter }}">
                    </div>

                    <div class="form-group">
                        <label class="form-label">状态</label>
                        <select class="form-select" name="status">
                            <option value="all" {% if status=="all" %}selected{% endif %}>全部状态</option>
                            <option value="pending" {% if status=="pending" %}selected{% endif %}>未安排</option>
                            <option value="assigned" {% if status=="assigned" %}selected{% endif %}>已安排</option>
                            <option value="cancelled" {% if status=="cancelled" %}selected{% endif %}>已取消</option>
                        </select>
                    </div>

                    <div class="form-group">
                        <label class="form-label">岗位</label>
                        <select class="form-select" name="role">
                            <option value="all" {% if role_filter=="all" %}selected{% endif %}>全部岗位</option>

                            {% for role in roles %}
                            <option value="{{ role }}" {% if role_filter==role %}selected{% endif %}>
                                {{ role }}
                            </option>
                            {% endfor %}
                        </select>
                    </div>

                    <div class="form-group">
                        <label class="form-label">&nbsp;</label>
                        <button class="btn-tool btn-blue btn-full" type="submit">
                            🔍 查询
                        </button>
                    </div>

                </div>

                <div class="quick-date-row">

                    <a class="btn-tool btn-orange"
                    href="/schedule/signups?date={{ today }}&status={{ status }}&role={{ role_filter }}">
                        今天
                    </a>

                    <a class="btn-tool btn-purple"
                    href="/schedule/signups?date={{ tomorrow }}&status={{ status }}&role={{ role_filter }}">
                        明天
                    </a>

                    <a class="btn-tool btn-gray"
                    href="/schedule/signups">
                        全部
                    </a>

                    <label class="history-check">
                        <input
                            type="checkbox"
                            name="show_history"
                            value="1"
                            {% if show_history %}checked{% endif %}>
                        显示过去记录
                    </label>

                </div>

            </form>

        </div>

        <div class="card">

            <h2 class="section-title">📄 报名记录</h2>

            <div class="table-wrap">

                <table class="signup-table">

                    <tr>
                        <th>日期</th>
                        <th>编号</th>
                        <th>姓名</th>
                        <th>岗位</th>
                        <th>时间</th>
                        <th>状态</th>
                        <th>系统安排</th>
                        <th>备注</th>
                        <th>操作</th>
                    </tr>

                    {% for r in rows %}
                    <tr class="status-{{ r.status }}">
                        <td>{{ r.signup_date }}</td>
                        <td>{{ r.volunteer_id }}</td>
                        <td><b>{{ r.name }}</b></td>
                        <td>{{ r.role }}</td>

                        <td>
                            {% if r.start_time or r.end_time %}
                                {{ r.start_time or "" }} ~ {{ r.end_time or "" }}
                            {% else %}
                                -
                            {% endif %}
                        </td>

                        <td>
                            {% if r.status == "pending" %}
                                <span class="badge badge-orange">未安排</span>
                            {% elif r.status == "assigned" %}
                                <span class="badge badge-green">已安排</span>
                            {% elif r.status == "cancelled" %}
                                <span class="badge badge-gray">已取消</span>
                            {% else %}
                                <span class="badge badge-blue">{{ r.status }}</span>
                            {% endif %}
                        </td>

                        <td class="assignment-text">
                            {% if r.final_assignment %}
                                {% for item in r.final_assignment %}
                                    <div>{{ item|safe }}</div>
                                {% endfor %}
                            {% else %}
                                -
                            {% endif %}
                        </td>

                        <td class="remark-text">
                            {{ r.remarks or "" }}
                        </td>

                        <td>
                            <div class="op-stack">

                                <form method="get" action="/schedule/signup_edit/{{ r.id }}">
                                    <button class="btn-tool btn-blue op-btn" type="submit">
                                        ✏️ 修改
                                    </button>
                                </form>

                                {% if r.status == "assigned" %}
                                <form method="get" action="/schedule/signup_place/{{ r.id }}">
                                    <button class="btn-tool btn-purple op-btn" type="submit">
                                        🔧 调整
                                    </button>
                                </form>
                                {% endif %}

                                {% if r.status != "cancelled" %}
                                <form
                                    method="post"
                                    action="/schedule/signup_cancel/{{ r.id }}"
                                    onsubmit="return confirm('确定取消这位义工的报名？');">
                                    <button class="btn-tool btn-red op-btn" type="submit">
                                        ❌ 取消
                                    </button>
                                </form>

                                {% else %}

                                <form
                                    method="post"
                                    action="/schedule/signup_restore/{{ r.id }}"
                                    onsubmit="return confirm('确定恢复报名？');">
                                    <button class="btn-tool btn-green op-btn" type="submit">
                                        ↩️ 恢复
                                    </button>
                                </form>

                                {% endif %}

                            </div>
                        </td>
                    </tr>
                    {% endfor %}

                </table>

            </div>

            {% if not rows %}
                <div class="empty-state">
                    <div class="empty-icon">📭</div>
                    <div class="empty-title">没有记录</div>
                    <div class="empty-text">请调整筛选条件后再查询。</div>
                </div>
            {% endif %}

        </div>

        <a href="/schedule/admin"
        class="btn-tool btn-blue floating-back">
            ⬅ 返回首页
        </a>

    </div>

    </body>
    </html>
    """,
        rows=rows,
        status=status,
        date_filter=date_filter,
        role_filter=role_filter,
        roles=ROLE_OPTIONS,
        summary=summary,
        role_summary=role_summary,
        today=today,
        tomorrow=tomorrow,
        show_history=show_history,
    )


LOGIN_HTML = """
<!doctype html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>负责人排班系统</title>

<link rel="stylesheet"
      href="{{ url_for('static', filename='css/toolbox.css') }}">

<style>
.login-wrap{
    max-width:620px;
    margin:60px auto;
    padding:20px;
}

.login-card{
    text-align:center;
}

.login-icon{
    font-size:64px;
    margin-bottom:10px;
}

.login-title{
    font-size:52px;
    font-weight:bold;
    color:#1e293b;
}

.login-subtitle{
    margin:15px 0 35px;
    color:#666;
    font-size:22px;
    line-height:1.8;
}

.login-input{
    margin-top:10px;
}

.login-input input{
    text-align:center;
    letter-spacing:6px;
    font-size:34px;
}

@media (max-width:700px){

    .login-title{
        font-size:40px;
    }

    .login-subtitle{
        font-size:20px;
    }

}
</style>

<script>
window.addEventListener("DOMContentLoaded",function(){

    document.querySelector("input[name='pin']").focus();

});
</script>

</head>

<body>

<div class="login-wrap">

    <div class="card login-card">

        <div class="login-icon">
            📅
        </div>

        <div class="login-title">
            负责人排班系统
        </div>

        <div class="login-subtitle">
            请输入负责人 PIN<br>
            进入排班管理系统。
        </div>

        <form method="post">

            <div class="form-group">

                <label class="form-label">
                    负责人 PIN
                </label>

                <div class="login-input">

                    <input
                        class="form-input"
                        type="password"
                        name="pin"
                        placeholder="请输入负责人 PIN"
                        inputmode="numeric"
                        autocomplete="new-password"
                        autocorrect="off"
                        autocapitalize="off"
                        spellcheck="false"
                        readonly
                        onfocus="this.removeAttribute('readonly');">

                </div>

            </div>

            <div class="btn-row">

                <button
                    class="btn-tool btn-green btn-full"
                    type="submit">

                    🔓 进入负责人系统

                </button>

            </div>

        </form>

    </div>

</div>

</body>
</html>
"""


SCHEDULE_HTML = """
<!doctype html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<link rel="manifest" href="/schedule-manifest.json">
<link rel="icon" href="/static/schedule_icon.png?v=1">
<title>负责人排班系统</title>

<link rel="stylesheet"
      href="{{ url_for('static', filename='css/toolbox.css') }}">


<script>
document.addEventListener("DOMContentLoaded", function () {

    function syncDates() {
        const mainDate = document.getElementById("main_date");

        const rawDate = document.getElementById("raw_single_date");
        const genDate = document.getElementById("generate_day_date");
        const fillDate = document.getElementById("fill_pending_date");
        const viewAssignedDate = document.getElementById("view_assigned_date");
        const overrideDate = document.getElementById("override_date");
        const prebookDate = document.getElementById("prebook_single_date");
        const publishDate = document.getElementById("publish_date");
        const noticeDate = document.getElementById("notice_date");

        if (!mainDate) return;

        if (rawDate) rawDate.value = mainDate.value;
        if (genDate) genDate.value = mainDate.value;
        if (fillDate) fillDate.value = mainDate.value;
        if (viewAssignedDate) viewAssignedDate.value = mainDate.value;
        if (overrideDate) overrideDate.value = mainDate.value;
        if (prebookDate) prebookDate.value = mainDate.value;
        if (publishDate) publishDate.value = mainDate.value;
        if (noticeDate) noticeDate.value = mainDate.value;
    }

    syncDates();

    const mainDate = document.getElementById("main_date");
    if (mainDate) {
        mainDate.addEventListener("change", syncDates);
    }
});
</script>

<script>
function showSpecial(id) {
    const boxes = document.querySelectorAll(".special-box");

    boxes.forEach(function(box) {
        box.style.display = "none";
    });

    const target = document.getElementById(id);

    if (target) {
        target.style.display = "block";
        setTimeout(function () {
            target.scrollIntoView({ behavior: "smooth", block: "start" });
        }, 80);
    }
}
</script>

<style>
.login-wrap{
    max-width:620px;
    margin:60px auto;
    padding:20px;
}

.login-card{
    text-align:center;
}

.login-icon{
    font-size:64px;
    margin-bottom:10px;
}

.login-title{
    font-size:52px;
    font-weight:bold;
    color:#1e293b;
}

.login-subtitle{
    margin:15px 0 35px;
    color:#666;
    font-size:22px;
    line-height:1.8;
}

.login-input{
    margin-top:10px;
}

.login-input input{
    text-align:center;
    letter-spacing:6px;
    font-size:34px;
}

@media (max-width:700px){
    .login-title{
        font-size:40px;
    }

    .login-subtitle{
        font-size:20px;
    }
}
</style>

</head>

<body>
<div class="box">

<div class="top-bar">
    <h1>📅 负责人排班系统</h1>

    <a class="btn-tool btn-gray" style="width:auto;min-width:160px;" href="/schedule/logout">
        🚪 退出登录
    </a>
</div>

<div class="card">
    <h2 class="section-title">📅 今日工作</h2>

    <div class="main-menu">

        <a class="btn-tool btn-blue"
           href="/schedule?mode=day&override_date={{ override_date }}">
            📅 当天安排
        </a>

        <a class="btn-tool btn-orange"
           href="/schedule/notice_center?date={{ override_date }}">
            📢 通知中心
        </a>

        <a class="btn-tool btn-purple"
           href="/schedule/signups">
            📋 报名管理
        </a>

    </div>
</div>

<div class="card">
    <h2 class="section-title">⚙️ 系统功能</h2>

    <div class="main-menu">

        <a class="btn-tool btn-gray"
           href="/schedule/settings">
            ⚙️ 排班设置
        </a>

        <a class="btn-tool btn-green"
           href="/reports">
            📚 资料中心
        </a>

        <a class="btn-tool btn-blue"
           href="/volunteer">
            👥 打开义工报名页面
        </a>

    </div>
</div>

{% if mode == "day" %}

<div class="card today-panel">
    <h2 class="section-title">📅 今日概况</h2>
    <div class="dashboard-note">负责人打开后，先确认日期、发布状态、报名人数和提醒。</div>

    <div class="date-card">
        <b>🗓 日期资料</b><br>
        阳历：{{ special_day_info.solar_date }}<br>
        农历：{{ special_day_info.lunar_text }}<br>

        {% if special_day_info.is_special %}
            <span class="text-red" style="font-weight:bold;">
                🔴 {{ special_day_info.special_names | join("、") }}
            </span><br>
            模板：{{ special_day_info.template_text }}<br>

            {% if special_day_info.setup_shifu %}
                🛕 需要设师父供台<br>
            {% endif %}

            {% if special_day_info.remove_next_day %}
                📦 明日需要收供台<br>
            {% endif %}
        {% else %}
            🟢 平时日
        {% endif %}

        {% if remove_info.need_remove_today_after_12 %}
            <br>⚠️ 今日中午后需要收昨日供台
        {% endif %}
    </div>

    <div class="date-status-row compact-date-row">

        <div class="date-picker-card compact-date-picker">

            <div class="status-title">
                📅 选择工作日期
            </div>

            <input
                type="date"
                id="main_date"
                class="date-input"
                value="{{ override_date }}"
                onchange="goToScheduleDate(this.value)"
            >

        </div>

    </div>

    <div class="kpi-grid">
        <div class="kpi-card">
            <div class="kpi-title">👥 已报名</div>
            <div class="kpi-value">{{ day_summary.total }}</div>
        </div>

        <div class="kpi-card">
            <div class="kpi-title">🟡 待安排</div>
            <div class="kpi-value">{{ day_summary.pending }}</div>
        </div>

        <div class="kpi-card">
            <div class="kpi-title">🟢 已安排</div>
            <div class="kpi-value">{{ day_summary.assigned }}</div>
        </div>

        <div class="kpi-card">
            <div class="kpi-title">⚪ 已取消</div>
            <div class="kpi-value">{{ day_summary.cancelled }}</div>
        </div>
    </div>

    <div class="info-box">
        <b>📢 缺人工提醒</b><br>

        {% if shortage_summary %}
            {% for line in shortage_summary %}
                {{ line }}<br>
            {% endfor %}
        {% else %}
            🟢 目前没有缺人工岗位
        {% endif %}
    </div>
</div>

<div class="card">
    <h2 class="section-title">🔔 今日提醒</h2>
    <div class="reminder-stack">
<div class="alert alert-error">
    <h3 style="margin-top:0;">🔴 未安排提醒</h3>
    今日：<b>{{ pending_counts.today }}</b> 人　
    明日：<b>{{ pending_counts.tomorrow }}</b> 人　
    本月：<b>{{ pending_counts.month }}</b> 人
</div>



        {% if republish_info and republish_info.need_republish %}
        <div class="alert alert-warning">
            🟠 正式值班表已有更新<br>
            最后更新时间：{{ republish_info.last_schedule_change }}<br>
            请到通知中心复制最新正式值班表 WhatsApp。
        </div>
        {% endif %}
<div class="alert alert-warning">
    <h2>🌸 供台报名提醒</h2>

    {% for item in supply_alerts %}
        <div style="margin-bottom:14px;">
            <b>
                {{ item.date.strftime("%m/%d") }}
                {% if item.type == "lunar_1_15" %}
                    初一 / 十五
                {% elif item.type == "buddhist_festival" %}
                    佛诞日
                {% endif %}
            </b><br>

            {% if item.names %}
                已报名：{{ "、".join(item.names) }}
            {% else %}
                <span style="color:#c0392b;">暂时没人报名</span>
            {% endif %}
        </div>
    {% endfor %}
</div>


    </div>
</div>

<div class="card">
    <h2 class="section-title">📝 今日操作</h2>

    <div class="action-grid">

        <form method="post" action="/schedule/publish">
            <input type="hidden" name="date" id="publish_date" value="{{ override_date }}">

            <button
                type="submit"
                class="btn-tool {% if is_published %}btn-gray{% else %}btn-green{% endif %}"
            >
                {% if is_published %}
                    ✅ 已发布
                    <span class="helper-text">正式值班表已生效，新报名会自动补排</span>
                {% else %}
                    🟢 发布正式值班表
                    <span class="helper-text">确认排班无误后再发布</span>
                {% endif %}
            </button>
        </form>

        {% if not is_published %}
            {% if not is_today_or_past %}
            <form method="post" action="/schedule/generate_day">
                <input type="hidden" name="date" id="generate_day_date" value="{{ override_date }}">
                <button type="submit" class="btn-tool btn-blue">
                    🔄 重新自动排班
                    <span class="helper-text">发布前可重复执行</span>
                </button>
            </form>
            {% else %}
            <form method="post" action="/schedule/fill_pending">
                <input type="hidden" name="date" id="fill_pending_date" value="{{ override_date }}">
                <button type="submit" class="btn-tool btn-blue">
                    ➕ 补入待安排
                    <span class="helper-text">今天或过去日期使用</span>
                </button>
            </form>
            {% endif %}
        {% endif %}

        <form method="post" action="/schedule/view_assigned">
            <input type="hidden" name="date" id="view_assigned_date" value="{{ override_date }}">
            <button type="submit" class="btn-tool btn-purple">
                👀 查看排班
                <span class="helper-text">查看 / 调整正式安排</span>
            </button>
        </form>

        <a href="/schedule/notice_center?date={{ override_date }}"
           class="btn-tool btn-orange">
            📢 通知中心
            <span class="helper-text">报名通知 / WhatsApp / 缺义工通知</span>
        </a>

        {% if is_published %}
        <a href="/schedule/attendance_status" class="btn-tool btn-blue">
            📋 签到情况
            <span class="helper-text">查看义工签到纪录</span>
        </a>
        {% endif %}

        <a href="/reports" class="btn-tool btn-gray">
            📚 资料中心
            <span class="helper-text">月报 / 年报 / 文件</span>
        </a>

    </div>
</div>



<div class="card">
    <h2 class="section-title">📋 报名录入</h2>
    <div class="dashboard-note">负责人只需要在这里加入报名：可以人工输入，也可以贴 WhatsApp 文字解析。</div>

    <div class="signup-unified">

        <div class="signup-block">
            <h3>➕ 负责人代报名</h3>

            <form method="post" action="/schedule/prebook_add">
                <input type="hidden" name="mode" value="day">
                <input type="hidden" name="single_date" id="prebook_single_date" value="{{ override_date }}">

                <div class="form-row-large">
                    <label>义工编号 / 姓名</label>

                    <div class="branch-search-row">
                        <button
                            type="button"
                            id="admin_branch_btn"
                            onclick="toggleAdminBranch()"
                            class="branch-toggle-btn">
                            CHE
                        </button>

                        <input type="hidden" id="admin_branch" name="branch" value="CHE">

                        <input
                            name="keyword"
                            id="admin_keyword"
                            placeholder="例如：108 / 姓名">
                    </div>

                    <div id="volunteer_lookup_result" class="lookup-text"></div>
                </div>

                {% for role in roles %}
                <label class="role-btn">

                    <input
                        type="checkbox"
                        name="roles"
                        value="{{ role }}">

                    {% if role == "值班" %}
                        👷 值班
                    {% elif role == "卫生" %}
                        🧹 卫生
                    {% elif role == "供台" %}
                        🛕 供台
                    {% else %}
                        {{ role }}
                    {% endif %}

                </label>
                {% endfor %}

                <div class="form-row-large">
                    <p>时间（只给值班用）</p>

                    <div class="time-grid">
                        <div>
                            <label>开始</label>
                            <select name="start_time" style="width:100%;">
                            {% for t in times %}
                                <option value="{{ t }}">{{ t }}</option>
                            {% endfor %}
                            </select>
                        </div>

                        <div>
                            <label>结束</label>
                            <select name="end_time" style="width:100%;">
                            {% for t in times %}
                                <option value="{{ t }}">{{ t }}</option>
                            {% endfor %}
                            </select>
                        </div>
                    </div>
                </div>

                <button type="submit" class="btn-tool btn-blue">➕ 加入报名</button>
            </form>
        </div>

        <div class="signup-divider"></div>

        <div class="signup-block">
            <h3>📥 WhatsApp 报名解析</h3>

            <form method="post" action="/schedule/parse_raw">
                <input type="hidden" name="single_date" id="raw_single_date" value="{{ override_date }}">

                <textarea
                    name="raw_signup"
                    rows="8"
                    placeholder="输入姓名 / 时间，例如：&#10;张三 10:00am-2:00pm&#10;李四 8:00am-12:00pm"></textarea>

                <br><br>

                <button type="submit" class="btn-tool btn-blue">📥 解析加入报名</button>
            </form>
        </div>

    </div>
</div>



    <div class="card">
        <h2 class="section-title">⚙️ 特殊设置</h2>

        <div class="quick-actions">
            <button type="button" class="btn-tool btn-gray" onclick="showSpecial('master-table')">
                🛕 师父供台
                <span class="helper-text">点击后自动跳到设置区</span>
            </button>

            <button type="button" class="btn-tool btn-gray" onclick="showSpecial('extra-buddha')">
                🌸 初一补人
            </button>

            <button type="button" class="btn-tool btn-gray" onclick="showSpecial('buddha-leave')">
                🔁 佛台请假
            </button>
        </div>

        <!-- 1. 师父供台 -->
        <div id="master-table" class="special-box" style="display:none;">
            <h3>🛕 师父供台设置</h3>

            {% if supply_signups %}
            <div style="
                background:#fff8e1;
                border:2px solid #f3d36b;
                border-radius:14px;
                padding:15px;
                margin:15px 0;
                font-size:22px;
            ">
                <b>🌸 已报名供台义工：</b><br>
                {% for p in supply_signups %}
                    {{ p.name }}{% if not loop.last %}、{% endif %}
                {% endfor %}
            </div>
            {% endif %}

            <form method="post" action="/schedule/save_day_flags">
                <input type="hidden" name="date" value="{{ override_date }}">

                <div class="info-box">
                    系统会根据初一、十五及佛诞日自动判断是否需要设供台。<br>
                    若前一天应收供台但未安排人员，系统会继续提醒收供台。
                </div>

                <h4>🪷 设师父供台人员</h4>
                <p style="color:#666;font-size:20px;">通常 1 至 2 位佛台组人员即可。</p>

                <select name="setup_person1">

                    <option value="">请选择</option>

                    {% for p in supply_signups %}
                        <option value="{{ p.name }}"
                            {% if loop.index == 1 %}selected{% endif %}>
                            {{ p.name }}
                        </option>
                    {% endfor %}

                    {% for n in buddha_names %}
                        <option value="{{ n }}"
                            {% if day_flags.setup_person1 == n %}selected{% endif %}>
                            {{ n }}
                        </option>
                    {% endfor %}

                </select>

                <select name="setup_person_2">
                    <option value="">请选择</option>

                    {% for p in supply_signups %}
                        <option value="{{ p.name }}"
                            {% if loop.index == 2 %}selected{% endif %}>
                            {{ p.name }}
                        </option>
                    {% endfor %}

                    {% for n in buddha_names %}
                        <option value="{{ n }}" {% if day_flags.setup_person_2 == n %}selected{% endif %}>
                            {{ n }}
                        </option>
                    {% endfor %}
                </select>

                <hr>

                <h4>📦 收师父供台人员</h4>
                <p style="color:#666;font-size:20px;">
                    通常 1 至 2 位佛台组人员。若没有安排人员，系统视为继续供奉。
                </p>

                <select name="remove_person_1">
                    <option value="">请选择</option>
                    {% for n in buddha_names %}
                    <option value="{{ n }}" {% if day_flags.remove_person_1 == n %}selected{% endif %}>
                        {{ n }}
                    </option>
                    {% endfor %}
                </select>

                <select name="remove_person_2">
                    <option value="">请选择</option>
                    {% for n in buddha_names %}
                    <option value="{{ n }}" {% if day_flags.remove_person_2 == n %}selected{% endif %}>
                        {{ n }}
                    </option>
                    {% endfor %}
                </select>

                <p>普通义工帮忙收供台：</p>
                <input
                    name="remove_extra_person"
                    value="{{ day_flags.remove_extra_person }}"
                    placeholder="例如：张三，可空"
                    style="width:100%;font-size:22px;"
                >

                <br><br>

                <button type="submit" class="btn-tool btn-blue">💾 保存供台人员</button>
            </form>
        </div>

        <!-- 2. 初一补人 -->
        <div id="extra-buddha" class="special-box" style="display:none;">
            <h3>🌸 初一整理佛台补人</h3>

            <div class="info-box">
                <b>🛕 当天负责整理佛台义工</b><br><br>

                {% if fixed_buddha_today %}
                    {% for n in fixed_buddha_today %}
                        🟢 {{ n }}<br>
                    {% endfor %}
                {% else %}
                    ⚪ 今天没有固定整理佛台义工
                {% endif %}
            </div>

            <form method="post" action="/schedule/save_day_flags">
                <input type="hidden" name="date" value="{{ override_date }}">

                <p style="font-size:20px;color:#666;">
                    仅初一需要。可增加第 3 位整理佛台义工。
                </p>

                <p>额外第 3 位整理佛台义工：</p>

                <select name="extra_buddha_person" style="width:100%;font-size:22px;">
                    <option value="">不需要增加</option>
                    {% for n in buddha_names %}
                    <option value="{{ n }}" {% if day_flags.extra_buddha_person == n %}selected{% endif %}>
                        {{ n }}
                    </option>
                    {% endfor %}
                </select>

                <br><br>

                <button type="submit" class="btn-tool btn-blue">💾 保存初一补人</button>
            </form>
        </div>

        <!-- 3. 佛台请假 -->
        <div id="buddha-leave" class="special-box" style="display:none;">
            <h3>🔁 佛台请假 / 换人</h3>

            <p>日期：{{ override_date }}</p>

            <form method="post" action="/schedule/override">
                <input type="hidden" name="date" id="override_date" value="{{ override_date }}">

                {% if fixed_buddha_today %}
                    {% for old_name in fixed_buddha_today %}
                        <div style="font-size:22px; margin:12px 0;">
                            {{ old_name }}
                            <input type="hidden" name="original_name" value="{{ old_name }}">

                            换成：
                            <select name="replacement_name">
                                <option value="">不换，保留原本</option>
                                <option value="__REMOVE__">请假，不找替补</option>

                                {% for n in buddha_names %}
                                <option value="{{ n }}">{{ n }}</option>
                                {% endfor %}
                            </select>
                        </div>
                    {% endfor %}

                    <button type="submit" class="btn-tool btn-blue">💾 保存佛台请假 / 换人</button>
                {% else %}
                    <p style="font-size:22px;">这一天没有固定佛台。</p>
                {% endif %}
            </form>
        </div>
    </div>

<div class="card">

    <h2 class="section-title">⚙️ 报名开放控制</h2>

    <table class="setting-table">

        <!-- 多日报名 -->

        <tr>
            <td style="padding:12px 0;">
                <b>📅 下个月多日报名</b><br>

                {% if multi_day_signup_open %}
                    <span style="color:green;font-weight:bold;">
                        🟢 已开放
                    </span>
                {% else %}
                    <span style="color:#999;font-weight:bold;">
                        ⚪ 已关闭
                    </span>
                {% endif %}
            </td>

            <td style="text-align:right;white-space:nowrap;">

                <form method="post"
                      action="/schedule/toggle_setting"
                      class="inline-form">

                    <input type="hidden"
                           name="key"
                           value="multi_day_signup_open">

                    <input type="hidden"
                           name="force"
                           value="true">

                    <button class="btn-tool btn-green mini-btn">
                        开启
                    </button>

                </form>

                <form method="post"
                      action="/schedule/toggle_setting"
                      class="inline-form">

                    <input type="hidden"
                           name="key"
                           value="multi_day_signup_open">

                    <input type="hidden"
                           name="force"
                           value="false">

                    <button class="btn-tool btn-red mini-btn">
                        关闭
                    </button>

                </form>

            </td>
        </tr>

        <tr>
            <td colspan="2">
                <hr>
            </td>
        </tr>

        <!-- 膳食 -->

        <tr>

            <td style="padding:12px 0;">

                <b>🍱 膳食组报名</b><br>

                {% if meal_signup_open %}
                    <span style="color:green;font-weight:bold;">
                        🟢 已开放
                    </span>
                {% else %}
                    <span style="color:#999;font-weight:bold;">
                        ⚪ 已关闭
                    </span>
                {% endif %}

            </td>

            <td style="text-align:right;white-space:nowrap;">

                <form method="post"
                      action="/schedule/toggle_setting"
                      class="inline-form">

                    <input type="hidden"
                           name="key"
                           value="meal_signup_open">

                    <input type="hidden"
                           name="force"
                           value="true">

                    <button class="btn-tool btn-green mini-btn">
                        开启
                    </button>

                </form>

                <form method="post"
                      action="/schedule/toggle_setting"
                      class="inline-form">

                    <input type="hidden"
                           name="key"
                           value="meal_signup_open">

                    <input type="hidden"
                           name="force"
                           value="false">

                    <button class="btn-tool btn-red mini-btn">
                        关闭
                    </button>

                </form>

            </td>

        </tr>

    </table>

</div>



{% else %}

<div class="card">
    <h2 class="section-title center-text">请选择要做的功能</h2>
    <p class="page-subtitle">建议先进入「当天安排」查看今日工作。</p>
</div>

{% endif %}

</div>
<script>
function toggleAdminBranch(){
    const btn = document.getElementById("admin_branch_btn");
    const branch = document.getElementById("admin_branch");

    if(branch.value === "CHE"){
        branch.value = "STW";
        btn.innerText = "STW";
        btn.style.background = "#dc3545";
    }else{
        branch.value = "CHE";
        btn.innerText = "CHE";
        btn.style.background = "#28a745";
    }
}
</script>
<script>
function goToScheduleDate(dateValue) {
    if (!dateValue) return;

    window.location.href =
        "/schedule?mode=day&override_date=" + encodeURIComponent(dateValue);
}
</script>
</body>
</html>
"""


DAY_OUTPUT_HTML = """
<!doctype html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>正式值班表</title>

<link rel="stylesheet"
      href="{{ url_for('static', filename='css/toolbox.css') }}">

<style>

.page-wrap{
    max-width:1100px;
    margin:auto;
    padding:25px;
}

.output-box{
    min-height:650px;
}

.output-box textarea{
    width:100%;
    height:650px;

    font-size:26px;
    line-height:1.7;

    padding:18px;

    border-radius:16px;
    border:2px solid #d6d6d6;

    box-sizing:border-box;

    resize:vertical;
}

</style>

<script>
function copySchedule(){

    const ta = document.getElementById("output");

    if(!ta){
        alert("❌ 找不到值班表内容");
        return;
    }

    ta.focus();
    ta.select();
    ta.setSelectionRange(0, 999999);

    const ok = document.execCommand("copy");

    if(ok){
        alert("✅ 已复制，可以贴去 WhatsApp 群");
    }else{
        alert("❌ 复制失败，请手动全选复制");
    }
}
</script>

</head>

<body>

<div class="page-wrap">

<div class="card">

    <h1 class="page-title">
        📋 正式值班表
    </h1>

    <p class="page-subtitle">
        可直接复制到 WhatsApp 发布。
    </p>

    <div class="btn-row">

        <a
            class="btn-tool btn-gray"
            href="/schedule?mode=day">

            ⬅ 返回负责人首页

        </a>

        <button
            type="button"
            class="btn-tool btn-green"
            onclick="copySchedule()">

            📋 一键复制

        </button>

    </div>

    <div class="output-box">

        <textarea
            id="output"
            readonly>{{ output }}</textarea>

    </div>

</div>

</div>

</body>
</html>
"""

MONTHLY_PREBOOK_HTML = """
<!doctype html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>月预报名表</title>

<link rel="stylesheet"
      href="{{ url_for('static', filename='css/toolbox.css') }}">

<style>
.page-wrap{
    max-width:1100px;
    margin:auto;
    padding:25px;
}

.output-box textarea{
    width:100%;
    height:650px;
    font-size:26px;
    line-height:1.7;
    padding:18px;
    border-radius:16px;
    border:2px solid #d6d6d6;
    box-sizing:border-box;
    resize:vertical;
}
</style>

<script>
function copyOutput(btn){
    const text = document.getElementById("output").value;
    navigator.clipboard.writeText(text);

    btn.innerText = "✅ 已复制";

    setTimeout(function(){
        btn.innerText = "📋 一键复制";
    }, 2000);
}
</script>

</head>

<body>

<div class="page-wrap">

<div class="card">

    <h1 class="page-title">📢 月预报名表</h1>

    <p class="page-subtitle">
        可直接复制到 WhatsApp 群发布。
    </p>

    <div class="btn-row">

        <a class="btn-tool btn-gray"
           href="/schedule?mode=prebook">
            ⬅ 返回月预报名
        </a>

        <button class="btn-tool btn-green"
                onclick="copyOutput(this)">
            📋 一键复制
        </button>

    </div>

    <div class="output-box">
        <textarea id="output" readonly>{{ output }}</textarea>
    </div>

</div>

</div>

</body>
</html>
"""